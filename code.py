#!/usr/bin/env python3
"""Build the self-contained Memory Observatory dashboard from the four workbooks.

Usage: python3 code.py [--data-dir DIR] [--output DIR] [--check]
Requires Python >=3.10 and openpyxl==3.1.5. No frontend dependencies or CDN.
All CSS, JavaScript, calculation rules, and README contents are in this file.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from hashlib import sha256
import json
import math
from pathlib import Path

from openpyxl import load_workbook

BUILD_VERSION = "1.0.0"
FILES = (
    "01_物料与项目基础数据_模拟数据.xlsx",
    "02_需求计划与库存供应_模拟数据.xlsx",
    "03_战备成本与价格风险_模拟数据.xlsx",
    "04_风险模型与处置闭环_模拟数据.xlsx",
)
KEYS = {
    "物料主数据": ["物料编码"], "项目主数据": ["项目编码"],
    "项目物料映射": ["项目编码", "物料编码", "生效日期"],
    "替代料与认证": ["原物料编码", "替代物料编码", "适用项目编码"],
    "组织责任映射": ["物料编码", "工厂"],
    "MDS需求版本": ["需求版本号", "需求周或月", "项目编码", "物料编码", "需求类别"],
    "实际消耗与出货": ["日期", "项目编码", "物料编码", "工厂"],
    "库存批次快照": ["快照日期", "仓库", "库存批次", "物料编码"],
    "采购订单明细": ["PO号", "PO行号"],
    "供应商交付与产能": ["统计周期", "供应商", "物料编码"],
    "库存策略参数": ["物料编码", "工厂", "策略生效日期"],
    "战略备货台账": ["战备决策编号"],
    "内存价格历史": ["价格日期", "物料编码", "价格类型", "地区", "价格来源"],
    "价格情景参数": ["情景编号", "物料编码", "模拟日期"],
    "库存财务成本": ["核算日期", "物料编码", "库存批次"],
    "供应质量风险": ["统计周期", "供应商", "物料编码"],
    "风险规则参数": ["规则编号", "规则版本"],
    "风险评估结果": ["评估日期", "物料编码", "项目编码", "库存批次"],
    "风险处置台账": ["风险事件编号"],
    "风险历史快照": ["快照日期", "风险事件编号"],
    "指标字典": ["指标名称", "规则版本"],
}


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def load_data(root: Path) -> dict:
    tables, sources, records = {}, [], {}
    for name in FILES:
        path = root / name
        book = load_workbook(path, read_only=True, data_only=True)
        source = {"name": name, "sha256": sha256(path.read_bytes()).hexdigest(), "sheets": [], "rows": 0}
        try:
            for sheet in book:
                iterator = sheet.iter_rows(values_only=True)
                columns = list(next(iterator))
                if not all(isinstance(c, str) and c.strip() for c in columns) or len(set(columns)) != len(columns):
                    raise ValueError(f"Invalid headers: {name}/{sheet.title}")
                rows = []
                for number, values in enumerate(iterator, 2):
                    if not any(v is not None for v in values):
                        continue
                    rows.append([number, *[clean(v) for v in values]])
                if sheet.title in tables:
                    raise ValueError(f"Duplicate sheet: {sheet.title}")
                tables[sheet.title] = {"columns": columns, "rows": rows, "source": name}
                records[sheet.title] = [dict(zip(columns, row[1:])) | {"__row": row[0]} for row in rows]
                source["sheets"].append(sheet.title)
                source["rows"] += len(rows)
        finally:
            book.close()
        sources.append(source)
    missing = set(KEYS) - set(tables)
    if missing or len(tables) != 21:
        raise ValueError(f"Expected 21 documented tables; missing: {sorted(missing)}")
    if any(not table["rows"] for table in tables.values()):
        raise ValueError("Empty input table")
    asof = max(str(r["快照日期"]) for r in records["库存批次快照"])
    issues = []

    def add(name, checked, failures, note, severity="error"):
        issues.append({"name": name, "checked": checked, "failed": len(failures),
                       "examples": failures[:30], "note": note, "severity": severity})

    for name, keys in KEYS.items():
        rows = records[name]
        counts = Counter(tuple(row[k] for k in keys) for row in rows)
        bad = [{"键": " / ".join(map(str, key)), "重复数": count} for key, count in counts.items() if count > 1]
        bad += [{"源行": row["__row"], "问题": "业务主键为空"} for row in rows if any(row[k] is None for k in keys)]
        add(f"主键 · {name}", len(rows), bad, " + ".join(keys))
    materials = {r["物料编码"] for r in records["物料主数据"]}
    projects = {r["项目编码"] for r in records["项目主数据"]}
    references = []
    for name, rows in records.items():
        for row in rows:
            for key in ("物料编码", "原物料编码", "替代物料编码", "项目编码", "关联项目编码", "适用项目编码"):
                if row.get(key) is not None and row[key] not in (projects if "项目" in key else materials):
                    references.append({"表": name, "源行": row["__row"], "字段": key, "值": row[key]})
    add("物料 / 项目外键", sum(len(v) for v in records.values()), references, "所有非空物料与项目引用必须可解析")
    po = {r["PO号"] for r in records["采购订单明细"]}
    lots = {r["库存批次"] for r in records["库存批次快照"]}
    events = {r["风险事件编号"] for r in records["风险处置台账"]}
    refs = []
    for name, fields in [("战略备货台账", [("对应PO号", po), ("对应库存批次", lots)]),
                         ("库存财务成本", [("库存批次", lots)]),
                         ("风险评估结果", [("库存批次", lots)]),
                         ("风险历史快照", [("风险事件编号", events)])]:
        for row in records[name]:
            for key, known in fields:
                if row.get(key) is not None and row[key] not in known:
                    refs.append({"表": name, "源行": row["__row"], "字段": key, "值": row[key]})
    add("PO / 批次 / 事件外键", len(po) + len(lots) + len(events), refs, "非空引用解析；空批次表示待到货，不伪造关联")
    bad = []
    for row in records["采购订单明细"]:
        if abs(row["订单数量"] - row["已收货数量"] - row["未交数量"]) > .01:
            bad.append({"PO": row["PO号"], "问题": "订单 ≠ 已收 + 未交"})
        if any(row[k] < 0 or row[k] > row["未交数量"] for k in ["可取消数量", "可延期数量", "运输中数量", "已确认待发数量"]):
            bad.append({"PO": row["PO号"], "问题": "未交子集越界"})
    add("PO 数量守恒", len(records["采购订单明细"]), bad, "运输中与待发不叠加为额外供应")
    bad = [{"批次": r["库存批次"]} for r in records["库存批次快照"]
           if abs(r["现有数量"] - r["可用数量"] - r["冻结数量"]) > .01]
    add("现有 / 可用 / 冻结勾稽", len(lots), bad, "现有 = 可用 + 冻结；已分配是否已扣仍需业务确认")
    bad = [{"批次": r["库存批次"], "差额": round(r["账面总成本"] - r["数量"] * r["账面单位成本"], 2)}
           for r in records["库存财务成本"] if abs(r["账面总成本"] - r["数量"] * r["账面单位成本"]) > .05]
    add("财务数量 × 单价", len(lots), bad, "允许 0.05 元四舍五入误差")
    bad = []
    for policy in records["库存策略参数"]:
        rows = [r for r in records["库存批次快照"] if r["物料编码"] == policy["物料编码"] and r["工厂"] == policy["工厂"]]
        snapshot_ss = sum(r["安全库存数量"] for r in rows)
        if abs(snapshot_ss - policy["安全库存数量"]) > .01:
            bad.append({"物料": policy["物料编码"], "快照安全库存合计": snapshot_ss, "策略安全库存": policy["安全库存数量"]})
    add("安全库存口径差异", len(materials), bad, "计算只按物料 + 工厂取最新有效策略一次，不累加批次安全库存", "warning")
    nulls = [{"表": name, "字段": col, "空值数": sum(row.get(col) is None for row in rows), "总行数": len(rows)}
             for name, rows in records.items() for col in tables[name]["columns"] if any(row.get(col) is None for row in rows)]
    prd = root / "看版prd.docx"
    return {"version": BUILD_VERSION, "asof": asof, "tables": tables, "sources": sources,
            "quality": issues, "nulls": nulls,
            "prd": {"name": prd.name, "sha256": sha256(prd.read_bytes()).hexdigest() if prd.exists() else None},
            "rowCount": sum(len(v) for v in records.values())}


HTML = r'''<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<link rel="icon" href="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 32 32'%3E%3Crect width='32' height='32' rx='9' fill='%231d1d1f'/%3E%3Crect x='9' y='9' width='14' height='14' rx='3' fill='none' stroke='white' stroke-width='2'/%3E%3Cpath d='M13 13h6v6h-6z' fill='none' stroke='white'/%3E%3C/svg%3E">
<meta name="color-scheme" content="light"><meta name="description" content="内存战略备货风险管理看板。四个工作簿，21 张数据表，离线交互与可追溯的风险分析。">
<title>Memory · 内存战略风险看板</title>
<style>
/* System fonts and native SVG: fully offline, no requests or external assets. */
:root{--bg:#f5f5f7;--panel:#fff;--ink:#1d1d1f;--muted:#86868b;--line:#e8e8ed;--blue:#0071e3;--blue-soft:#edf5ff;--green:#248a54;--red:#d64045;--amber:#b97d12;--radius:20px;--shadow:0 3px 18px #1d1d1f04}
*{box-sizing:border-box}body{margin:0;font-family:-apple-system,BlinkMacSystemFont,"SF Pro Display","PingFang SC","Microsoft YaHei",sans-serif;background:var(--bg);color:var(--ink);font-size:13px;-webkit-font-smoothing:antialiased}button,input,select,textarea{font:inherit}button,a,summary,input,select{touch-action:manipulation}button{cursor:pointer}a{color:var(--blue);text-decoration:none}button{border:0;background:transparent;color:inherit}button:disabled{opacity:.4;cursor:not-allowed}button:focus-visible,a:focus-visible,summary:focus-visible{outline:3px solid #0071e355;outline-offset:3px}input:focus,select:focus,textarea:focus{outline:2px solid #0071e34d;outline-offset:1px}input,select,textarea{border:1px solid var(--line);border-radius:9px;background:#fff;color:var(--ink);padding:9px 11px;min-width:0}input[type=range]{accent-color:var(--blue);width:100%;padding:0}input[type=checkbox]{accent-color:var(--blue);width:15px;height:15px}textarea{resize:vertical;min-height:90px}svg{display:block}h1,h2,h3,p{margin:0}h1{font-size:34px;line-height:1.2;font-weight:650;letter-spacing:-1.5px}h2{font-size:17px;font-weight:600;letter-spacing:-.4px}h3{font-size:14px;font-weight:600}small,.small{font-size:11px}.muted{color:var(--muted)}.mono{font-variant-numeric:tabular-nums;font-feature-settings:'tnum';letter-spacing:-.3px}.positive{color:var(--green)}.negative{color:var(--red)}.amber{color:var(--amber)}.blue{color:var(--blue)}.sr-only{position:absolute;clip:rect(0,0,0,0);width:1px;height:1px;overflow:hidden}.app{display:flex;min-height:100vh}.sidebar{width:216px;position:fixed;inset:0 auto 0 0;background:#fbfbfddd;border-right:1px solid var(--line);padding:29px 16px;display:flex;flex-direction:column;z-index:30;backdrop-filter:blur(20px)}.brand{display:flex;align-items:center;gap:10px;padding:0 12px 30px;font-size:19px;font-weight:650;letter-spacing:-.7px}.brandmark{background:var(--ink);color:#fff;width:31px;height:31px;border-radius:10px;display:grid;place-items:center}.brandmark svg{width:19px}.nav-caption{font-size:10px;letter-spacing:1.7px;color:#a1a1a6;padding:8px 13px 12px}.nav{display:grid;gap:5px}.nav button{display:flex;align-items:center;gap:11px;text-align:left;padding:12px 13px;border-radius:10px;color:#66666c;font-weight:500;transition:background .18s}.nav button svg{width:18px;height:18px;stroke-width:1.6}.nav button:hover{background:#ededf1}.nav button.active{background:var(--blue-soft);color:var(--blue);font-weight:600}.nav .count{margin-left:auto;background:#0071e315;padding:2px 6px;border-radius:5px;font-size:10px}.sidebar-bottom{margin-top:auto;padding:24px 12px 0;line-height:1.8}.live-dot{display:inline-block;width:6px;height:6px;background:var(--green);border-radius:50%;margin-right:6px}.avatar{width:30px;height:30px;background:#ececf0;border-radius:50%;display:grid;place-items:center;font-size:11px;color:#777}.workspace{margin-left:216px;width:calc(100% - 216px);min-width:0}.topbar{height:70px;background:#f5f5f7df;backdrop-filter:blur(20px);position:sticky;top:0;z-index:20;border-bottom:1px solid var(--line);display:flex;align-items:center;justify-content:space-between;padding:0 34px;gap:12px}.breadcrumb{color:var(--muted);display:flex;align-items:center;gap:12px;font-size:12px}.breadcrumb b{color:var(--ink);font-weight:500}.actions{display:flex;gap:8px;align-items:center}.btn{display:inline-flex;gap:7px;align-items:center;justify-content:center;padding:9px 13px;border:1px solid var(--line);border-radius:9px;background:#fff;font-size:12px;white-space:nowrap}.btn:hover{background:#f5f5f7}.btn svg{width:15px;height:15px}.btn.primary{background:var(--blue);border-color:var(--blue);color:#fff}.btn.primary:hover{background:#0077ed}.btn.ghost{border-color:transparent;background:transparent;color:var(--blue)}.iconbtn{width:32px;height:32px;padding:7px;border-radius:9px;display:grid;place-items:center}.iconbtn:hover{background:#ededf1}.iconbtn svg{width:18px;height:18px}.main{max-width:1680px;margin:auto;padding:32px 34px 24px}.pagehead{display:flex;justify-content:space-between;align-items:flex-end;gap:20px;margin-bottom:25px}.eyebrow{font-size:10px;font-weight:600;letter-spacing:2px;color:var(--muted);margin-bottom:10px}.subtitle{color:#808087;font-size:12px;margin-top:11px;line-height:1.8}.page-meta{text-align:right;color:var(--muted);font-size:11px;line-height:1.9}.tag{display:inline-flex;gap:5px;align-items:center;padding:4px 8px;border-radius:6px;background:#f0f0f3;font-size:10px;color:#73737b;white-space:nowrap}.tag.blue{background:var(--blue-soft);color:var(--blue)}.tag.red{background:#fff0f0;color:var(--red)}.tag.green{background:#edf8f0;color:var(--green)}.tag.amber{background:#fcf5e8;color:var(--amber)}.filters{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:12px}.filter{position:relative}.filter summary{list-style:none;background:#fff;border:1px solid var(--line);border-radius:9px;padding:9px 12px;cursor:pointer;display:flex;gap:15px;align-items:center;font-size:12px;white-space:nowrap}.filter summary::-webkit-details-marker{display:none}.filter summary:after{content:'⌄';color:#aaa}.filter[open]{z-index:26}.filter[open] summary{border-color:#99c6f3}.popover{position:absolute;top:calc(100% + 7px);left:0;background:#fff;border:1px solid var(--line);box-shadow:0 12px 40px #0002;border-radius:13px;padding:8px;min-width:200px;max-width:310px;max-height:320px;overflow:auto;z-index:28}.popover label{display:flex;gap:9px;align-items:center;padding:8px;border-radius:7px;font-size:12px;cursor:pointer}.popover label:hover{background:var(--bg)}.popover input[type=search]{width:100%;margin-bottom:5px}.filter.active summary{color:var(--blue);background:#f4f9ff;border-color:#cce3fb}.search{display:flex;align-items:center;position:relative;flex:1;min-width:165px;max-width:280px}.search svg{position:absolute;left:11px;width:15px;height:15px;color:var(--muted)}.search input{padding-left:34px;width:100%;font-size:12px}.chips{display:flex;gap:6px;flex-wrap:wrap;min-height:27px;align-items:center;margin-bottom:16px}.chip{border-radius:6px;padding:4px 8px;color:var(--blue);background:#e8f1fc;font-size:10px;display:flex;gap:7px;align-items:center}.view-note{font-size:10px;color:var(--muted);margin-left:auto}.panel{background:var(--panel);border:1px solid var(--line);border-radius:var(--radius);box-shadow:var(--shadow);min-width:0;overflow:hidden}.panel-pad{padding:22px 24px}.panelhead{display:flex;justify-content:space-between;align-items:center;gap:12px;margin-bottom:18px}.panelhead p{color:var(--muted);font-size:11px;margin-top:5px;line-height:1.7}.kpis{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px;margin-bottom:18px}.kpi{padding:19px 21px;position:relative;overflow:hidden}.kpi-label{font-size:11px;color:#717179;display:flex;justify-content:space-between;align-items:center;gap:6px}.kpi-value{font-size:32px;line-height:1.3;font-weight:600;letter-spacing:-1.3px;margin:10px 0 8px;white-space:nowrap}.kpi-value span{font-size:13px;font-weight:500;margin-left:4px;letter-spacing:0;color:#929298}.kpi-note{font-size:10px;color:var(--muted);display:flex;align-items:center;gap:6px;line-height:1.6}.kpi.accent{background:linear-gradient(130deg,#fff,#f1f7ff);border-color:#dceafa}.grid-main{display:grid;grid-template-columns:minmax(0,1.8fr) minmax(270px,1fr);gap:18px;margin-bottom:18px}.grid-two{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:18px;margin-bottom:18px}.grid-three{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px;margin-bottom:18px}.stack{display:grid;gap:18px;align-content:start}.seg{padding:3px;background:#f3f3f6;border-radius:8px;display:inline-flex;gap:2px;align-items:center;flex-shrink:0}.seg button{padding:5px 10px;font-size:10px;border-radius:6px;color:#7b7b81;white-space:nowrap}.seg button.active{background:#fff;box-shadow:0 1px 4px #0001;color:var(--ink)}.chart{width:100%;position:relative}.chart svg{width:100%;height:auto;overflow:visible}.chart text{font-family:inherit;font-size:10px;fill:#98989e}.chart .gridline{stroke:#ededf1;stroke-dasharray:3 4}.chart [data-tip]{cursor:crosshair}.chart [data-material]{cursor:pointer}.legend{display:flex;gap:15px;font-size:10px;color:#929298;align-items:center;flex-wrap:wrap;margin-bottom:12px}.dot{width:6px;height:6px;display:inline-block;border-radius:50%;margin-right:5px}.legend .line{width:15px;height:2px;display:inline-block;vertical-align:middle;margin-right:5px}.note{color:#8d8d94;font-size:10px;line-height:1.8;margin-top:11px}.notice{border:1px solid #e0eaf6;background:#f2f7fd;border-radius:12px;padding:12px 15px;color:#5f7895;font-size:11px;line-height:1.8;margin-bottom:18px;display:flex;align-items:flex-start;gap:9px}.notice.warning{border-color:#efe5ce;background:#fcf8ef;color:#9b793f}.notice svg{flex-shrink:0;width:16px;height:16px;margin-top:2px}.risk-item{display:flex;align-items:center;gap:12px;padding:13px 0;border-bottom:1px solid #f1f1f4}.risk-item:last-child{border-bottom:0;padding-bottom:0}.risk-icon{width:32px;height:32px;display:grid;place-items:center;border-radius:10px;background:#fff1f0;color:var(--red);flex-shrink:0}.risk-icon.amber{background:#fbf4e7;color:var(--amber)}.risk-icon.blue{background:var(--blue-soft);color:var(--blue)}.risk-icon svg{width:16px;height:16px}.risk-item .body{flex:1}.risk-item .body p{font-size:11px;color:var(--muted);margin-top:5px;line-height:1.6}.risk-item strong{font-size:21px;font-weight:550}.table-panel{margin-bottom:18px}.table-head{padding:20px 24px 15px;display:flex;justify-content:space-between;align-items:center;gap:12px;flex-wrap:wrap}.table-wrap{overflow:auto;max-height:580px}.table-wrap table{width:100%;border-collapse:separate;border-spacing:0;font-size:11px;white-space:nowrap}.table-wrap th{text-align:left;background:#fafafb;color:#8b8b93;font-size:10px;font-weight:500;padding:12px 16px;position:sticky;top:0;z-index:3;border-bottom:1px solid var(--line)}.table-wrap th button{text-align:left;font:inherit;width:100%;white-space:nowrap;display:flex;gap:7px;align-items:center}.table-wrap td{padding:14px 16px;border-bottom:1px solid #f1f1f4;max-width:280px;overflow:hidden;text-overflow:ellipsis}.table-wrap tr:last-child td{border-bottom:0}.table-wrap tbody tr:hover td{background:#fafbfd}.table-wrap th:first-child,.table-wrap td:first-child{position:sticky;left:0;background:#fff;z-index:2}.table-wrap th:first-child{z-index:4;background:#fafafb}.table-wrap .material{font-size:11px;display:flex;flex-direction:column;gap:4px;align-items:flex-start;text-align:left}.material b{font-weight:600;color:var(--blue)}.material small{color:var(--muted);font-size:9px;max-width:160px;overflow:hidden;text-overflow:ellipsis}.table-footer{padding:12px 22px;border-top:1px solid var(--line);display:flex;align-items:center;justify-content:space-between;gap:8px;color:var(--muted);font-size:10px}.table-tools{display:flex;gap:7px;align-items:center}.table-tools input{max-width:160px;font-size:11px;padding:7px 10px}.table-tools .btn{font-size:10px;padding:7px 9px}.table-tools .filter summary{padding:7px 10px;font-size:10px}.bar-track{height:5px;width:85px;background:#f0f0f4;border-radius:5px;overflow:hidden;margin-top:5px}.bar-track i{display:block;height:100%;border-radius:5px;background:var(--blue)}.tinybar{width:65px;height:24px}.pill{padding:4px 7px;border-radius:5px;background:#f2f2f5;color:#8a8a90;font-size:10px}.pill.high{color:var(--red);background:#fff1f0}.pill.medium{color:var(--amber);background:#fff8e8}.pill.low{color:var(--green);background:#eef7f1}.empty{padding:40px 20px;text-align:center;color:#97979f;font-size:12px;line-height:2}.empty strong{display:block;color:#6c6c73;font-size:14px;font-weight:500}.foot{display:flex;justify-content:space-between;gap:12px;font-size:10px;color:#a1a1a8;line-height:1.8;padding:10px 2px}.mobile-toggle{display:none}.tooltip{position:fixed;pointer-events:none;z-index:100;background:#1d1d1ff2;color:#fff;border-radius:10px;padding:10px 13px;font-size:11px;line-height:1.9;max-width:310px;box-shadow:0 5px 20px #0002;display:none;white-space:pre-line}.toast{position:fixed;bottom:25px;left:50%;transform:translate(-50%,15px);opacity:0;z-index:120;background:#242426;color:white;padding:12px 20px;border-radius:12px;font-size:12px;box-shadow:0 8px 30px #0002;pointer-events:none;transition:all .2s}.toast.show{opacity:1;transform:translate(-50%,0)}.overlay{position:fixed;inset:0;background:#0003;backdrop-filter:blur(3px);z-index:50;display:flex;justify-content:flex-end}.drawer{width:min(810px,94vw);height:100%;overflow:auto;background:var(--bg);box-shadow:-10px 0 60px #0001;padding:30px}.drawer-head{display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:22px}.drawer h1{font-size:27px}.drawer .kpis{grid-template-columns:repeat(3,minmax(0,1fr))}.drawer .kpi-value{font-size:23px}.drawer .panel{margin-bottom:15px}.drawer .grid-two{gap:12px}.drawer .table-wrap{max-height:380px}.definition{display:grid;grid-template-columns:130px 1fr;gap:12px;font-size:11px;line-height:1.8}.definition dt{color:var(--muted)}.definition dd{margin:0}.divider{border:0;border-top:1px solid var(--line);margin:20px 0}.formula{font-family:ui-monospace,SFMono-Regular,monospace;font-size:11px;padding:14px 16px;border-radius:10px;background:#f6f7fa;line-height:2;color:#64748b;white-space:normal}.fieldgrid{display:grid;grid-template-columns:1fr 1fr;gap:15px}.field{display:flex;flex-direction:column;gap:7px;font-size:11px;color:#7b7b83}.field input,.field select,.field textarea{width:100%;color:var(--ink)}.field.full{grid-column:1/-1}.dialog{border:1px solid var(--line);border-radius:20px;padding:25px;width:min(570px,94vw);max-height:90vh;overflow:auto;color:var(--ink);box-shadow:0 25px 100px #0002}.dialog::backdrop{background:#0004;backdrop-filter:blur(5px)}.dialog-head{display:flex;justify-content:space-between;align-items:center;margin-bottom:20px}.dialog-foot{display:flex;justify-content:flex-end;gap:8px;margin-top:20px}.slider-value{font-size:48px;letter-spacing:-2px;font-weight:600;margin:15px 0}.scenario-buttons{display:flex;gap:7px;flex-wrap:wrap;margin-top:12px}.scenario-buttons button{font-size:10px;padding:6px 9px;border:1px solid var(--line);border-radius:6px}.scenario-buttons button.active{color:var(--blue);border-color:#bcdbfa;background:#f1f7ff}.metricrow{display:flex;justify-content:space-between;align-items:center;padding:13px 0;border-bottom:1px solid var(--line);font-size:11px;color:var(--muted)}.metricrow strong{font-size:17px;color:var(--ink);font-weight:550}.sourcecards{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;margin-bottom:18px}.sourcecard{padding:18px}.sourcecard .num{font-size:28px;margin:12px 0 6px;font-weight:550;letter-spacing:-1px}.sourcecard p{font-size:10px;color:var(--muted);line-height:1.8;word-break:break-all}.audit{border-left:1px solid #d7e3f3;margin:10px 0 15px 5px;padding-left:18px}.audit-item{position:relative;padding:0 0 20px}.audit-item:before{content:'';position:absolute;width:7px;height:7px;background:#8cb8e7;border-radius:50%;left:-22px;top:4px}.audit-item p{font-size:11px;margin-top:6px;line-height:1.8;color:var(--muted)}.settingrow{display:flex;justify-content:space-between;gap:15px;align-items:center;padding:15px 0;border-bottom:1px solid var(--line);font-size:12px}.settingrow p{font-size:10px;color:var(--muted);margin-top:6px;line-height:1.8}.settingrow input[type=number]{width:90px}.settingrow select{max-width:190px}.settingrow:last-child{border:0}.progress-line{height:7px;display:flex;border-radius:8px;overflow:hidden;background:#f1f1f4;margin:17px 0}.progress-line i{height:100%}.link{color:var(--blue);font-size:11px;cursor:pointer;padding:0}.statusselect{max-width:150px;font-size:11px}.paper{color:var(--muted);font-size:11px;line-height:1.9}.data-select{max-width:270px;font-size:12px}.compare-card{border-radius:13px;padding:18px;background:#f7f8fa}.compare-card .label{color:var(--muted);font-size:11px}.compare-card strong{display:block;font-size:25px;font-weight:550;margin-top:8px}.page-fade{animation:enter .25s ease-out}@keyframes enter{from{opacity:.25;transform:translateY(5px)}to{opacity:1;transform:translateY(0)}}
@media(min-width:1650px){.kpis{grid-template-columns:repeat(6,minmax(0,1fr))}.kpi-value{font-size:30px}}
@media(max-width:1200px){.sidebar{width:190px}.workspace{margin-left:190px;width:calc(100% - 190px)}.main{padding:25px 24px}.topbar{padding:0 24px}.grid-main{grid-template-columns:minmax(0,1.6fr) minmax(240px,1fr)}.panel-pad{padding:20px}.kpi{padding:17px}.kpi-value{font-size:28px}.sourcecards{grid-template-columns:repeat(2,minmax(0,1fr))}}
@media(max-width:980px){.sidebar{width:74px;padding:27px 10px}.brand{padding:0 11px 30px}.brand span,.nav-caption,.nav button span,.sidebar-bottom{display:none}.nav button{padding:13px;justify-content:center}.nav button svg{width:21px;height:21px}.workspace{margin-left:74px;width:calc(100% - 74px)}.grid-main{grid-template-columns:1fr}.grid-two{grid-template-columns:1fr}.kpis{grid-template-columns:repeat(3,minmax(0,1fr))}.page-meta{display:none}}
@media(max-width:640px){.sidebar{display:none}.sidebar.open{display:flex;width:216px;box-shadow:0 0 0 100vmax #0005}.sidebar.open .brand span,.sidebar.open .nav-caption,.sidebar.open .nav button span,.sidebar.open .sidebar-bottom{display:block}.sidebar.open .nav button{justify-content:flex-start}.sidebar.open .brand{padding-left:13px}.workspace{width:100%;margin-left:0}.mobile-toggle{display:grid}.topbar{height:58px;padding:0 15px}.breadcrumb>span{display:none}.topbar .btn{padding:7px 9px;font-size:10px}.main{padding:25px 16px 15px}h1{font-size:28px}.kpis{grid-template-columns:repeat(2,minmax(0,1fr));gap:9px}.kpi{padding:15px 12px}.kpi-value{font-size:25px;letter-spacing:-1px}.kpi-label{font-size:10px}.kpi-note{font-size:9px}.panel-pad{padding:18px}.table-head{padding:17px;gap:12px}.table-head h2{font-size:15px}.table-tools{flex-wrap:wrap}.table-tools input{max-width:130px}.grid-three{grid-template-columns:1fr}.sourcecards{gap:9px}.sourcecard{padding:14px}.pagehead{margin-bottom:20px}.filters{gap:6px}.filter summary{font-size:11px;padding:8px 10px}.search{max-width:none;flex-basis:100%}.view-note{margin:3px 0}.chips{margin-bottom:12px}.foot{flex-direction:column;gap:3px}.drawer{padding:20px 14px;width:97vw}.drawer .kpis{grid-template-columns:repeat(2,minmax(0,1fr))}.fieldgrid{grid-template-columns:1fr}.field.full{grid-column:auto}.definition{grid-template-columns:95px 1fr}.panelhead{align-items:flex-start}.seg button{padding:5px 7px}.subtitle{font-size:11px}.kpi-value span{font-size:11px}.notice{font-size:10px}}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}html{scroll-behavior:auto}}
.popover label[hidden]{display:none}
@media(max-width:640px){.topbar #snapshotTag{display:none}.breadcrumb b{white-space:nowrap}.chart text{font-size:18px}}
@media print{.sidebar,.topbar,.filters,.chips,.table-tools,.table-footer,.actions,.overlay,.tooltip,.toast,.btn,.seg,.popover{display:none!important}.workspace{margin:0;width:100%}.main{padding:0}.panel{break-inside:avoid;border-color:#ddd;box-shadow:none}.table-wrap{max-height:none;overflow:visible}.table-wrap th,.table-wrap td{position:static!important}body{background:#fff}.grid-main,.grid-two{display:block}.panel{margin-bottom:15px}.foot{display:block}.kpis{grid-template-columns:repeat(3,1fr)}}
</style>
</head>
<body>
<div class="app">
<aside class="sidebar" id="sidebar"><div class="brand"><div class="brandmark" id="brandmark"></div><span>Memory<span style="color:#a1a1a6">.</span></span></div><div class="nav-caption">风险管理工作台</div><nav class="nav" id="nav" aria-label="主导航"></nav><div class="sidebar-bottom"><div><span class="live-dot"></span><span style="font-size:11px">离线数据已就绪</span></div><p class="muted small" style="margin-top:5px">模拟数据 · V1.0<br>仅本地演示与分析</p><hr class="divider"><div style="display:flex;align-items:center;gap:9px"><div class="avatar">M</div><div style="font-size:10px">本地分析工作台<br><span class="muted">Local workspace</span></div></div></div></aside>
<div class="workspace"><header class="topbar"><div style="display:flex;align-items:center;gap:9px"><button class="iconbtn mobile-toggle" data-action="menu" aria-label="展开导航">☰</button><div class="breadcrumb"><span>战略库存</span><span>/</span><b id="crumb">管理总览</b></div></div><div class="actions"><span class="tag" id="snapshotTag"></span><button class="btn" data-action="method">指标口径</button><button class="btn" data-action="print" id="printButton"></button></div></header><main class="main" id="main"></main></div>
</div>
<div class="tooltip" id="tooltip" role="tooltip"></div><div class="toast" id="toast" role="status" aria-live="polite"></div>
<div id="drawerRoot"></div><dialog id="dialog" class="dialog"></dialog>
<script type="application/json" id="embedded-data">__DASHBOARD_DATA__</script>
<script>
'use strict';
// SECTION: DATA / STATE / PURE CALCULATIONS
const DATA=JSON.parse(document.getElementById('embedded-data').textContent);
const T=Object.fromEntries(Object.entries(DATA.tables).map(([name,t])=>[name,t.rows.map(row=>Object.fromEntries([['__row',row[0]],...t.columns.map((c,i)=>[c,row[i+1]])]))]));
const ASOF=DATA.asof, MS=86400000, RULE_VERSION='demo-1.0', STORAGE='memory-dashboard-v1-'+ASOF;
const $=s=>document.querySelector(s), $$=s=>[...document.querySelectorAll(s)];
const esc=v=>String(v??'—').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const attr=esc, num=v=>Number.isFinite(Number(v))?Number(v):0;
const sum=(a,f=x=>x)=>a.reduce((s,x)=>s+num(f(x)),0), uniq=a=>[...new Set(a.filter(x=>x!==null&&x!==undefined&&x!==''))];
const group=(a,k)=>a.reduce((o,r)=>((o[r[k]]??=[]).push(r),o),{});
const utc=d=>new Date(String(d).slice(0,10)+'T00:00:00Z').getTime();
const iso=t=>new Date(t).toISOString().slice(0,10);
const monday=d=>{const t=utc(d),day=new Date(t).getUTCDay();return iso(t-((day+6)%7)*MS)};
const days=(a,b)=>Math.round((utc(a)-utc(b))/MS);
const short=d=>d?String(d).slice(5).replace('-','/'):'—';
const fmt=(v,d=0)=>v===null||v===undefined||!Number.isFinite(Number(v))?'—':Number(v).toLocaleString('zh-CN',{maximumFractionDigits:d,minimumFractionDigits:d});
const money=v=>Math.abs(v)>=1e8?'¥'+fmt(v/1e8,2)+' 亿':'¥'+fmt(v/1e4,2)+' 万';
const qty=v=>Math.abs(v)>=1e4?fmt(v/1e4,2)+' 万':fmt(v);
const pct=v=>Number.isFinite(v)?fmt(v*100,1)+'%':'—';
const nowCN=()=>new Date().toLocaleString('zh-CN',{timeZone:'Asia/Shanghai',hour12:false});
const MATERIALS=T['物料主数据'], MM=Object.fromEntries(MATERIALS.map(r=>[r['物料编码'],r]));
const PM=Object.fromEntries(T['项目主数据'].map(r=>[r['项目编码'],r]));
const ORG=Object.fromEntries(T['组织责任映射'].map(r=>[r['物料编码'],r]));
const versions=uniq(T['MDS需求版本'].filter(r=>r['版本发布日期']<=ASOF).map(r=>r['需求版本号'])).sort();
const idx={};for(const name of Object.keys(T))idx[name]=group(T[name],'物料编码');
const weeks=uniq(T['MDS需求版本'].filter(r=>r['需求版本号']===versions.at(-1)).map(r=>monday(r['需求周或月']))).sort();
const defaultPolicy={allocatedIncluded:false,materialDemand:true,excludeCancelled:true,coverage:21,otd:85,priceLimit:1000000};
let storageWarning='';
function readLocal(){try{const v=JSON.parse(localStorage.getItem(STORAGE)||'{}');return v&&typeof v==='object'&&!Array.isArray(v)?v:{}}catch(e){storageWarning='浏览器本地存储不可用，本次操作仅在当前页面保留。';return {}}}
let local=readLocal();
local={events:{},created:[],audit:[],views:[],...local};
if(!Array.isArray(local.created))local.created=[];if(!Array.isArray(local.audit))local.audit=[];if(!Array.isArray(local.views))local.views=[];if(!local.events||typeof local.events!=='object'||Array.isArray(local.events))local.events={};
const policy={...defaultPolicy};
const S={page:'overview',filters:{bu:[],factory:[],family:[],project:[],material:[],supplier:[],category:[],direction:[],level:[]},query:'',version:versions.at(-1),horizon:13,scenario:-20,chartClass:'DRAM',chartMode:'balance',demandScenario:'base',eventStatus:'全部',source:'物料主数据',tables:{},drawer:null};
let models=[],filtered=[],tableRegistry={},toastTimer,returnFocus;
const projectLinks=group(T['项目物料映射'],'物料编码');
function roll(initial,demands,arrivals,safety){let balance=initial;return demands.map((d,i)=>{const opening=balance;balance+=num(arrivals[i])-num(d);return {week:weeks[i],opening,arrival:num(arrivals[i]),demand:num(d),balance,shortage:Math.max(0,-balance),safetyGap:Math.max(0,safety-balance)}})}
function poAmounts(row){const open=Math.max(0,num(row['未交数量'])),eligible=!!row['最后变更日期']&&row['最后变更日期']>=ASOF,cancel=eligible?Math.min(open,Math.max(0,num(row['可取消数量']))):0,defer=eligible?Math.min(open,Math.max(0,num(row['可延期数量']))):0;return {open,cancel,defer,locked:open-cancel,price:num(row['采购单价']),expired:!!row['最后变更日期']&&!eligible,unknown:!row['最后变更日期']}}
function exposure(batches,price){return sum(batches,r=>num(r['数量'])*Math.max(0,num(r['账面单位成本'])-price))}
function demandRows(id,version){return (idx['MDS需求版本'][id]||[]).filter(r=>r['需求版本号']===version&&(!S.filters.project.length||S.filters.project.includes(r['项目编码'])))}
function demandVector(id,version,scenario=true,applyCancellation=true){const out=weeks.map(()=>0);for(const r of demandRows(id,version)){const j=weeks.indexOf(monday(r['需求周或月']));if(j<0)continue;const p=PM[r['项目编码']]||{};if(applyCancellation&&policy.excludeCancelled&&p['项目状态']==='取消'&&p['取消或暂停日期']&&r['需求周或月']>=p['取消或暂停日期'])continue;const bom=(projectLinks[id]||[]).find(x=>x['项目编码']===r['项目编码']&&x['生效日期']<=r['需求周或月']&&(!x['失效日期']||x['失效日期']>=r['需求周或月']));let value=num(r['需求数量'])*(policy.materialDemand?1:num(bom?.['单机用量']));if(scenario&&r['是否锁定']!=='是')value*=S.demandScenario==='stress'?1.2:S.demandScenario==='low'?.8:1;out[j]+=value}return out}
function buildModels(){
 const vi=versions.indexOf(S.version),previous=versions[vi-1];
 models=MATERIALS.map(mat=>{
  const id=mat['物料编码'],org=ORG[id]||{},lots=(idx['库存批次快照'][id]||[]).filter(r=>r['快照日期']===ASOF),finance=(idx['库存财务成本'][id]||[]).filter(r=>r['核算日期']===ASOF),pos=idx['采购订单明细'][id]||[];
  const policies=(idx['库存策略参数'][id]||[]).filter(r=>r['策略生效日期']<=ASOF).sort((a,b)=>b['策略生效日期'].localeCompare(a['策略生效日期']));
  const selectedPolicies=[...new Map(policies.map(r=>[r['工厂'],policies.find(x=>x['工厂']===r['工厂'])])).values()];
  const safety=sum(selectedPolicies,r=>r['安全库存数量']),targetDays=num(selectedPolicies[0]?.['目标库存天数'])||35;
  const onhand=sum(lots,r=>r['现有数量']),available=sum(lots,r=>r['可用数量']),allocated=sum(lots,r=>r['已分配数量']),initial=available-(policy.allocatedIncluded?0:allocated);
  const cost=sum(finance,r=>r['账面总成本']),unitCost=onhand?cost/onhand:0;
  const demand=demandVector(id,S.version),prior=previous?demandVector(id,previous):null,arrivals=weeks.map(()=>0),overdue=[];
  for(const p of pos){const w=monday(p['承诺到货日期']),j=weeks.indexOf(w);if(p['承诺到货日期']<ASOF){if(num(p['未交数量'])>0)overdue.push(p)}else if(j>=0)arrivals[j]+=num(p['未交数量'])}
  const rolling=roll(initial,demand,arrivals,safety),future=demand.slice(0,4),daily=sum(future)/28,coverage=daily>0?initial/daily:null;
  // Integrate weekly buckets fractionally for a non-multiple-of-seven target.
  let targetDemand=0;for(let i=0;i<demand.length;i++)targetDemand+=demand[i]*Math.max(0,Math.min(1,(targetDays-i*7)/7));
  const target=targetDemand+safety,excess=Math.max(0,initial-target),locked=sum(pos,p=>poAmounts(p).locked),cancelAmount=sum(pos,p=>poAmounts(p).cancel*num(p['采购单价'])),deferAmount=sum(pos,p=>poAmounts(p).defer*num(p['采购单价']));
  const excessCommitted=Math.max(0,initial+locked-target),strategies=idx['战略备货台账'][id]||[];
  const stratByLot={};for(const r of strategies){if(r['对应库存批次'])stratByLot[r['对应库存批次']]=(stratByLot[r['对应库存批次']]||0)+num(r['已到货数量'])}
  const stratQty=sum(lots,r=>Math.min(num(r['现有数量']),stratByLot[r['库存批次']]||0));
  const strategicExcess=onhand?excess*unitCost*stratQty/onhand:0;
  const prices=(idx['内存价格历史'][id]||[]).filter(r=>r['价格日期']<=ASOF&&r['币种']==='CNY'&&r['报价单位']==='元/颗').sort((a,b)=>b['价格日期'].localeCompare(a['价格日期']));
  const market=num(prices[0]?.['市场价格']),scenarioPrice=market*(1+S.scenario/100),priceExposure=exposure(finance,scenarioPrice),baselineExposure=exposure(finance,market);
  const excessExposure=onhand?priceExposure*Math.min(excess,onhand)/onhand:0;
  const supply=(idx['供应商交付与产能'][id]||[]).filter(r=>r['统计周期']<=ASOF).sort((a,b)=>b['统计周期'].localeCompare(a['统计周期']))[0];
  const quality=(idx['供应质量风险'][id]||[]).filter(r=>r['统计周期']<=ASOF).sort((a,b)=>b['统计周期'].localeCompare(a['统计周期']))[0];
  const seed=(idx['风险评估结果'][id]||[])[0]||{},sameWindow=sum(demand.slice(0,13)),rawCurrent=sum(demandVector(id,S.version,false,false).slice(0,13)),previousDemand=previous?sum(demandVector(id,previous,false,false).slice(0,13)):null,decline=previousDemand>0?(previousDemand-rawCurrent)/previousDemand:null;
  const first=rolling.findIndex(r=>r.balance<0),firstSafety=rolling.findIndex(r=>r.safetyGap>0),maxGap=Math.max(0,...rolling.slice(0,S.horizon).map(r=>r.shortage));
  const relatedProjects=(projectLinks[id]||[]).map(r=>PM[r['项目编码']]).filter(Boolean),cancelled=relatedProjects.filter(p=>p['项目状态']==='取消'&&p['取消或暂停日期']<=ASOF);
  const flags=[];if(first>=0&&first<4)flags.push({id:'SUP-01',level:'严重',text:'未来4周实物缺料',direction:'供应'});if(coverage!==null&&coverage<policy.coverage&&daily>0)flags.push({id:'SUP-02',level:'高',text:'库存覆盖低于阈值',direction:'供应'});if(supply&&(num(supply['OTD'])<policy.otd/100||num(supply['交期波动天数'])>7))flags.push({id:'SUP-03',level:'高',text:'供应商交付风险',direction:'供应'});
  const historyDemands=versions.filter(v=>v<=S.version).slice(-3).map(v=>sum(demandVector(id,v,false,false).slice(0,13)));
  if(coverage!==null&&coverage>num(selectedPolicies[0]?.['战略覆盖上限天数']||100)&&historyDemands.length===3&&historyDemands[2]<historyDemands[1]&&historyDemands[1]<historyDemands[0])flags.push({id:'INV-01',level:'高',text:'超覆盖且连续两版下调',direction:'库存'});
  if(cancelled.length&&(onhand>0||locked>0))flags.push({id:'INV-02',level:'严重',text:'取消项目存在共享料库存或PO',direction:'库存'});
  if(lots.some(r=>num(r['库龄天数'])>90)&&sameWindow===0)flags.push({id:'INV-03',level:'高',text:'长库龄且13周无需求',direction:'库存'});
  if(exposure(finance,market*.8)>policy.priceLimit)flags.push({id:'PRC-01',level:'高',text:'跌价20%敞口超阈值',direction:'库存'});
  if(pos.some(p=>p['最后变更日期']&&days(p['最后变更日期'],ASOF)>=0&&days(p['最后变更日期'],ASOF)<=7&&num(p['未交数量'])>0))flags.push({id:'PO-01',level:'高',text:'PO干预窗口即将到期',direction:'库存'});
  const level=flags.some(r=>r.level==='严重')?'严重':flags.length?'高':'低';
  const hold=sum(finance,r=>num(r['账面总成本'])*(num(r['年化资金成本率'])+num(r['仓储费率']))*90/365);
  return {id,mat,org,lots,finance,pos,policies:selectedPolicies,safety,targetDays,onhand,available,allocated,initial,cost,unitCost,demand,prior,arrivals,rolling,coverage,daily,target,excess,excessCommitted,locked,cancelAmount,deferAmount,strategies,stratQty,strategicExcess,prices,market,scenarioPrice,priceExposure,baselineExposure,excessExposure,supply,quality,seed,decline,first,firstSafety,maxGap,relatedProjects,cancelled,flags,level,hold,overdue};
 });
}
function matches(m){const f=S.filters;return (!f.bu.length||f.bu.includes(m.org['业务单元']))&&(!f.factory.length||f.factory.includes(m.org['工厂']))&&(!f.family.length||f.family.includes(m.org['产品族']))&&(!f.project.length||m.relatedProjects.some(p=>f.project.includes(p['项目编码'])))&&(!f.material.length||f.material.includes(m.id))&&(!f.supplier.length||f.supplier.includes(m.mat['供应商']))&&(!f.category.length||f.category.includes(m.mat['DRAM_NAND分类']))&&(!f.direction.length||m.flags.some(x=>f.direction.includes(x.direction)))&&(!f.level.length||f.level.includes(m.level))&&(!S.query||[m.id,m.mat['物料名称'],m.org['风险负责人'],...m.relatedProjects.map(p=>p['项目名称'])].some(v=>String(v||'').toLowerCase().includes(S.query.toLowerCase())))}
function getEvents(){return [...T['风险处置台账'],...local.created].map(r=>({...r,...(local.events[r['风险事件编号']]||{})}))}
function currentEvents(){const ids=new Set(filtered.map(m=>m.id));return getEvents().filter(r=>ids.has(r['物料编码'])&&(!S.filters.project.length||S.filters.project.includes(r['项目编码'])))}
function overdueEvent(r){return r['处置状态']!=='已关闭'&&r['截止日期']<ASOF}
function persist(){try{localStorage.setItem(STORAGE,JSON.stringify(local));return true}catch(e){storageWarning='本地存储失败，请立即导出工作区备份。';toast(storageWarning);return false}}
function toast(s){$('#toast').textContent=s;$('#toast').classList.add('show');clearTimeout(toastTimer);toastTimer=setTimeout(()=>$('#toast').classList.remove('show'),3000)}
const ICONS={chip:'<rect x="5" y="5" width="14" height="14" rx="3"/><path d="M9 9h6v6H9zM9 2v3m6-3v3M9 19v3m6-3v3M2 9h3m-3 6h3m14-6h3m-3 6h3"/>',overview:'<rect x="3" y="3" width="7" height="7" rx="1.5"/><rect x="14" y="3" width="7" height="7" rx="1.5"/><rect x="3" y="14" width="7" height="7" rx="1.5"/><rect x="14" y="14" width="7" height="7" rx="1.5"/>',supply:'<path d="M3 16l5-6 4 3 8-10M15 3h5v5M3 21h18"/>',inventory:'<path d="M3 7l9-5 9 5-9 5zM3 7v10l9 5 9-5V7M12 12v10M7.5 4.5l9 5"/>',price:'<path d="M3 17l5-5 4 3 8-10M16 5h4v4M3 21h18"/><path d="M5 3v4m-2-2h4"/>',po:'<rect x="5" y="3" width="14" height="18" rx="2"/><path d="M9 8h6M9 12h6m-6 4h3"/>',risks:'<path d="M12 3l10 18H2zM12 9v5m0 3v.1"/>',data:'<ellipse cx="12" cy="5" rx="8" ry="3"/><path d="M4 5v7c0 4 16 4 16 0V5M4 12v7c0 4 16 4 16 0v-7"/>',download:'<path d="M12 3v12m-5-5l5 5 5-5M4 15v5h16v-5"/>',search:'<circle cx="10.5" cy="10.5" r="6.5"/><path d="M16 16l5 5"/>',info:'<circle cx="12" cy="12" r="9"/><path d="M12 11v6m0-10v.2"/>',arrow:'<path d="M5 12h14m-5-5l5 5-5 5"/>',close:'<path d="M6 6l12 12M18 6L6 18"/>',check:'<path d="M5 12l4 4L20 5"/>',clock:'<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/>',print:'<path d="M7 8V3h10v5M7 17H4V9h16v8h-3M7 14h10v7H7z"/>',plus:'<path d="M12 5v14M5 12h14"/>'};
function icon(n){return `<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">${ICONS[n]||ICONS.info}</svg>`}
function badge(v){const c=['高','严重','逾期'].includes(v)?'high':['中','待确认','已缓解','待复核','暂停'].includes(v)?'medium':['低','已关闭','已认证'].includes(v)?'low':'';return `<span class="pill ${c}">${esc(v)}</span>`}
function materialCell(id){return `<button class="material" data-material="${attr(id)}"><b>${esc(id)} ↗</b><small>${esc(MM[id]?.['物料名称']||'新增物料')}</small></button>`}
function tip(s){return `data-tip="${attr(s)}"`}
function kpi(label,value,unit,note='',tone='',info=''){return `<article class="panel kpi ${tone}" ${info?tip(info):''}><div class="kpi-label">${label}${info?'<span class="muted">ⓘ</span>':''}</div><div class="kpi-value mono">${value}<span>${unit}</span></div><div class="kpi-note">${note}</div></article>`}
function moneyKpi(label,value,note='',tone='',info=''){return kpi(label,fmt(value/(Math.abs(value)>=1e8?1e8:1e4),2),Math.abs(value)>=1e8?'亿元':'万元',note,tone,info)}
function empty(text='试试放宽筛选条件，或重置当前视图。'){return `<div class="empty"><strong>暂无匹配数据</strong>${esc(text)}</div>`}
function panel(title,subtitle,body,tools=''){return `<section class="panel panel-pad"><div class="panelhead"><div><h2>${title}</h2>${subtitle?`<p>${subtitle}</p>`:''}</div>${tools}</div>${body}</section>`}
const PAGES={overview:['管理总览','OVERVIEW','风险，尽在掌握。','连接需求、库存与价格，让每一次战略备货都有据可依。'],supply:['供应保障','SUPPLY ASSURANCE','提前看见每一个缺口。','按自然周滚动，区分实物缺料与安全库存缺口。'],inventory:['战略库存','STRATEGIC INVENTORY','让库存保持从容。','追溯战备决策，识别需求变化、长库龄与超额占用。'],price:['价格与情景','PRICE & SCENARIOS','为价格波动，做好准备。','从批次账面成本出发，探索每一种价格情景。'],po:['PO 干预','PURCHASE ORDERS','在窗口关闭之前。','逐笔审视未交承诺、可调整数量与最后变更日期。'],risks:['风险处置','RISK WORKSPACE','把发现，变成行动。','以事件为中心，追踪责任、措施与复核证据。'],data:['数据质量与配置','DATA & GOVERNANCE','每一个数字，都有来处。','21 张原始数据表、字段质量与可解释的计算口径。']};
// SECTION: REUSABLE TABLE / CHART COMPONENTS
function table(id,title,cols,rows,subtitle='',size=8){
 const st=S.tables[id]??={sort:null,dir:1,page:1,search:'',hidden:[]};
 let list=rows.filter(r=>!st.search||cols.some(c=>String(c.value(r)??'').toLowerCase().includes(st.search.toLowerCase())));
 const sc=cols.find(c=>c.key===st.sort);if(sc)list=list.map((r,i)=>({r,i})).sort((a,b)=>{const x=sc.value(a.r),y=sc.value(b.r);let n=(typeof x==='number'&&typeof y==='number')?x-y:String(x??'').localeCompare(String(y??''),'zh-CN',{numeric:true});return n*st.dir||a.i-b.i}).map(x=>x.r);
 const pages=Math.max(1,Math.ceil(list.length/size));st.page=Math.min(Math.max(1,st.page),pages);const shown=list.slice((st.page-1)*size,st.page*size),visible=cols.filter(c=>!st.hidden.includes(c.key));
 tableRegistry[id]={title,cols,rows:list};
 const tools=`<div class="table-tools"><input aria-label="搜索${attr(title)}" placeholder="搜索此表…" data-table-search="${id}" value="${attr(st.search)}"><details class="filter"><summary aria-label="显示列">列</summary><div class="popover">${cols.map((c,i)=>`<label><input type="checkbox" data-column="${attr(c.key)}" data-table="${id}" ${st.hidden.includes(c.key)?'':'checked'} ${i===0?'disabled':''}>${esc(c.label)}</label>`).join('')}</div></details><button class="btn" data-action="export" data-id="${id}">${icon('download')}CSV</button></div>`;
 return `<section class="panel table-panel" id="table-${id}"><div class="table-head"><div><h2>${title}</h2>${subtitle?`<p class="note" style="margin-top:5px">${subtitle}</p>`:''}</div>${tools}</div>${list.length?`<div class="table-wrap"><table><thead><tr>${visible.map(c=>`<th aria-sort="${st.sort===c.key?(st.dir===1?'ascending':'descending'):'none'}"><button data-action="sort" data-id="${id}" data-key="${attr(c.key)}">${esc(c.label)} <span>${st.sort===c.key?(st.dir===1?'↑':'↓'):'↕'}</span></button></th>`).join('')}</tr></thead><tbody>${shown.map(r=>`<tr>${visible.map(c=>`<td>${c.render?c.render(r):esc(c.value(r)??'—')}</td>`).join('')}</tr>`).join('')}</tbody></table></div>`:empty()}<div class="table-footer"><span>共 ${fmt(list.length)} 条${st.search?' · 已筛选':''} · 每页 ${size} 条</span><div style="display:flex;gap:12px;align-items:center"><button class="iconbtn" aria-label="上一页" data-action="paginate" data-id="${id}" data-step="-1" ${st.page===1?'disabled':''}>‹</button><span>${st.page} / ${pages}</span><button class="iconbtn" aria-label="下一页" data-action="paginate" data-id="${id}" data-step="1" ${st.page===pages?'disabled':''}>›</button></div></div></section>`;
}
const col=(key,label,value,render)=>({key,label,value,render});
const materialCol=col('id','物料 / 规格',m=>m.id,m=>materialCell(m.id));
function rawTable(name,id='source'){const cols=[col('__row','Excel 行',r=>r.__row),...DATA.tables[name].columns.map(c=>col(c,c,r=>r[c],r=>c==='物料编码'?materialCell(r[c]):esc(r[c]??'—')))];return table(id,esc(name),cols,T[name],`${esc(DATA.tables[name].source)} · 保留源行号 · 原始模拟数据`,10)}
function lineChart(series,labels,{height=250,unit='万颗',scale=1e4,zero=true,aria='趋势图',marker=null}={}){
 const W=720,H=height,L=56,R=20,TOP=22,B=35,n=labels.length;
 const all=series.flatMap(s=>s.values).filter(Number.isFinite);if(!n||!all.length)return empty('当前范围没有可绘制的数据。');
 let lo=Math.min(...all,zero?0:Infinity),hi=Math.max(...all,0);if(hi===lo){hi+=1;lo=Math.min(0,lo-1)}const spread=hi-lo;hi+=spread*.12;lo-=lo<0?spread*.08:0;
 const x=i=>L+(n===1?.5:i/(n-1))*(W-L-R),y=v=>TOP+(hi-v)/(hi-lo)*(H-TOP-B);
 const grid=Array.from({length:5},(_,i)=>{const v=lo+(hi-lo)*i/4;return `<line class="gridline" x1="${L}" x2="${W-R}" y1="${y(v)}" y2="${y(v)}"/><text x="${L-10}" y="${y(v)+3}" text-anchor="end">${fmt(v/scale,Math.abs(v/scale)<10?1:0)}</text>`}).join('');
 const stamp=labels.map((l,i)=>i===0||i===n-1||i%Math.max(1,Math.ceil(n/6))===0?`<text x="${x(i)}" y="${H-10}" text-anchor="middle">${esc(short(l))}</text>`:'').join('');
 const paths=series.map((s,k)=>{const points=s.values.map((v,i)=>`${x(i)},${y(v)}`).join(' '),last=s.values.length-1;return `${s.area?`<polygon points="${x(0)},${y(Math.max(lo,0))} ${points} ${x(last)},${y(Math.max(lo,0))}" fill="${s.color}" opacity=".055"/>`:''}<polyline points="${points}" fill="none" stroke="${s.color}" stroke-width="${s.width||2.3}" ${s.dash?'stroke-dasharray="5 5"':''} stroke-linejoin="round" stroke-linecap="round"/>${s.values.map((v,i)=>`<circle cx="${x(i)}" cy="${y(v)}" r="${n>18?5:7}" fill="transparent" ${tip(`${labels[i]} · ${s.name}\n${fmt(v/scale,2)} ${unit}\n数据快照 ${ASOF}`)}/>`).join('')}`}).join('');
 const zeroLine=lo<0?`<line x1="${L}" x2="${W-R}" y1="${y(0)}" y2="${y(0)}" stroke="#c7c7cc" stroke-width="1"/>`:'';
 return `<div class="legend">${series.map(s=>`<span><i class="line" style="background:${s.color}"></i>${esc(s.name)}</span>`).join('')}<span style="margin-left:auto">${esc(unit)}</span></div><div class="chart"><svg viewBox="0 0 ${W} ${H}" role="img" aria-label="${attr(aria)}">${grid}${zeroLine}${stamp}${paths}</svg></div>`;
}
function rollingChart(list=filtered,mode=S.chartMode,horizon=S.horizon){const within=list.filter(m=>m.mat['DRAM_NAND分类']===S.chartClass);if(!within.length)return empty(`当前筛选没有 ${S.chartClass} 物料，可切换 DRAM / NAND。`);const labels=weeks.slice(0,horizon),values=k=>labels.map((_,i)=>sum(within,m=>m.rolling[i][k]));let series=mode==='gap'?[{name:'实物缺料',color:'#d64045',values:values('shortage'),area:true},{name:'安全库存缺口',color:'#cf9b4d',values:values('safetyGap'),dash:true}]:mode==='excess'?[{name:'期末超目标库存',color:'#b58a43',values:labels.map((_,i)=>sum(within,m=>Math.max(0,m.rolling[i].balance-m.target))),area:true}]:[{name:'期末实物库存',color:'#0071e3',values:values('balance'),area:true},{name:'安全库存',color:'#c7c7cc',values:labels.map(()=>sum(within,m=>m.safety)),dash:true}];return lineChart(series,labels,{aria:`${S.chartClass} ${horizon}周库存滚动`})}
function chartTools(){return `<div class="seg">${['DRAM','NAND'].map(v=>`<button data-action="chart-class" data-value="${v}" class="${S.chartClass===v?'active':''}">${v}</button>`).join('')}</div>`}
function bubbleChart(){const rows=filtered.filter(m=>m.coverage!==null&&m.decline!==null);if(!rows.length)return empty('需要前一版 MDS 与非零需求，才能比较同窗变化。');const W=640,H=280,L=58,R=25,TOP=22,B=43,xmin=Math.min(-.05,...rows.map(m=>m.decline)),xmax=Math.max(.5,...rows.map(m=>m.decline)),ymax=Math.max(6,...rows.map(m=>Math.min(m.coverage/30,18))),x=v=>L+(v-xmin)/(xmax-xmin)*(W-L-R),y=v=>H-B-Math.min(v,ymax)/ymax*(H-TOP-B),max=Math.max(1,...rows.map(m=>m.excess*m.unitCost));return `<div class="chart"><svg viewBox="0 0 ${W} ${H}" role="img" aria-label="需求下调率与库存覆盖月数风险矩阵"><rect x="${x(.15)}" y="${TOP}" width="${W-R-x(.15)}" height="${Math.max(0,y(3.33)-TOP)}" fill="#fdf8f3" rx="5"/>${[0,1,2,3].map(i=>{const v=ymax*i/3;return `<line class="gridline" x1="${L}" x2="${W-R}" y1="${y(v)}" y2="${y(v)}"/><text x="${L-10}" y="${y(v)+4}" text-anchor="end">${fmt(v,1)}</text>`}).join('')}${[0,.25,.5,.75,1].map(i=>{const v=xmin+(xmax-xmin)*i;return `<text x="${x(v)}" y="${H-23}" text-anchor="middle">${pct(v)}</text>`}).join('')}<text x="${L}" y="11">库存覆盖 / 月（最多绘制18个月）</text><text x="${W-R}" y="${H-3}" text-anchor="end">同窗需求下调率 →</text>${rows.map(m=>`<circle cx="${x(m.decline)}" cy="${y(Math.min(18,m.coverage/30))}" r="${5+Math.sqrt(m.excess*m.unitCost/max)*15}" fill="${m.level==='严重'?'#d99570':m.level==='高'?'#b2c8e3':'#b7c8bd'}" opacity=".7" stroke="#fff" stroke-width="1.5" data-material="${m.id}" tabindex="0" role="button" aria-label="查看${m.id}" ${tip(`${m.id} · ${m.mat['物料名称']}\n需求下调 ${pct(m.decline)} · 覆盖 ${fmt(m.coverage,1)} 天\n现货超额成本 ${money(m.excess*m.unitCost)}\n点击查看计算明细`)}/>`).join('')}</svg></div><p class="note">气泡大小 = 现货超额成本；右上角表示需求下调与库存高覆盖并存。点击气泡下钻。</p>`}
function miniBars(values,color='#0071e3'){const max=Math.max(1,...values.map(Math.abs));return `<svg class="tinybar" viewBox="0 0 65 24" aria-hidden="true">${values.map((v,i)=>`<rect x="${i*5}" y="${22-Math.abs(v)/max*20}" width="3" height="${Math.max(1,Math.abs(v)/max*20)}" rx="1" fill="${v<0?'#dc7d7d':color}" opacity="${.35+i/values.length*.6}"/>`).join('')}</svg>`}
// SECTION: GLOBAL FILTERS / SHELL
const FILTERS={bu:['业务单元',()=>uniq(T['组织责任映射'].map(r=>r['业务单元']))],factory:['工厂',()=>uniq(T['组织责任映射'].map(r=>r['工厂']))],family:['产品族',()=>uniq(T['组织责任映射'].map(r=>r['产品族']))],project:['项目',()=>T['项目主数据'].map(r=>r['项目编码'])],material:['物料',()=>MATERIALS.map(r=>r['物料编码'])],supplier:['供应商',()=>uniq(MATERIALS.map(r=>r['供应商']))],category:['品类',()=>['DRAM','NAND']],direction:['风险方向',()=>['供应','库存']],level:['风险等级',()=>['严重','高','低']]};
function filterName(k,v){return k==='project'?`${v} · ${PM[v]?.['项目名称']||''}`:k==='material'?`${v} · ${MM[v]?.['物料名称']||''}`:v}
function filterHTML(){return `<div class="filters">${Object.entries(FILTERS).map(([key,[label,options]])=>`<details class="filter ${S.filters[key].length?'active':''}" data-filter-details="${key}"><summary>${label}${S.filters[key].length?` <span>${S.filters[key].length}</span>`:''}</summary><div class="popover"><input type="search" placeholder="搜索${label}…" aria-label="搜索${label}选项" data-option-search="${key}">${options().map(v=>`<label data-option-label="${attr(filterName(key,v))}"><input type="checkbox" data-filter="${key}" value="${attr(v)}" ${S.filters[key].includes(v)?'checked':''}>${esc(filterName(key,v))}</label>`).join('')}</div></details>`).join('')}<label class="search">${icon('search')}<input id="globalSearch" aria-label="搜索物料、项目或责任人" placeholder="搜索物料、项目或责任人" value="${attr(S.query)}"></label><button class="btn ghost" data-action="reset">重置</button></div><div class="chips">${Object.entries(S.filters).flatMap(([k,vs])=>vs.map(v=>`<button class="chip" data-action="remove-filter" data-key="${k}" data-value="${attr(v)}">${esc(v)} <span>×</span></button>`)).join('')}${S.query?`<span class="chip">搜索：${esc(S.query)}</span>`:''}<span class="view-note">${filtered.length} / ${MATERIALS.length} 个物料 · ${S.version} · 快照固定 ${ASOF}</span><button class="link" data-action="save-view">保存视图</button>${local.views.length?'<button class="link" data-action="load-view">读取视图</button>':''}</div>`}
function render({focus=null,position=null}={}){
 buildModels();filtered=models.filter(matches);tableRegistry={};
 $('#nav').innerHTML=Object.entries(PAGES).map(([id,v])=>`<button data-action="nav" data-page="${id}" class="${S.page===id?'active':''}" ${S.page===id?'aria-current="page"':''}>${icon(id)}<span>${v[0]}</span>${id==='risks'?`<span class="count">${getEvents().filter(r=>r['处置状态']!=='已关闭').length}</span>`:''}</button>`).join('');
 $('#brandmark').innerHTML=icon('chip');$('#printButton').innerHTML=icon('print')+'<span>打印 / PDF</span>';$('#snapshotTag').textContent='数据快照 '+ASOF;$('#crumb').textContent=PAGES[S.page][0];
 const p=PAGES[S.page];$('#main').innerHTML=`<div class="pagehead"><div><div class="eyebrow">${p[1]}</div><h1>${p[2]}</h1><p class="subtitle">${p[3]}</p></div><div class="page-meta"><span class="tag blue">模拟数据 · 本地交互版</span><br>业务基准 ${ASOF} · CNY<br>规则版本 ${RULE_VERSION}</div></div>${filterHTML()}${storageWarning?`<div class="notice warning">${esc(storageWarning)}</div>`:''}${S.filters.project.length?'<div class="notice">ⓘ 项目筛选仅限定需求与事件，并选择关联物料。库存及 PO 为这些物料的完整共享池，尚未分摊为项目专属库存。</div>':''}<div class="page-fade">${({overview:overviewPage,supply:supplyPage,inventory:inventoryPage,price:pricePage,po:poPage,risks:risksPage,data:dataPage}[S.page])()}</div><footer class="foot"><span>Memory. 内存战略备货风险管理 · 合成演示数据，非会计减值或采购执行依据</span><span>四个工作簿 · 21 张表 · 本地处理，无外部请求</span></footer>`;
 if(focus){const e=$(focus);if(e){e.focus({preventScroll:true});if(position!==null&&e.setSelectionRange)e.setSelectionRange(position,position)}}
}
function demandControls(){return `<div class="filters"><label class="muted small">MDS 版本 <select id="versionSelect" aria-label="MDS版本">${versions.map(v=>`<option ${S.version===v?'selected':''}>${v}</option>`).join('')}</select></label><label class="muted small"> 需求情景 <select id="demandScenario" aria-label="需求情景"><option value="base" ${S.demandScenario==='base'?'selected':''}>基准 · 原始需求</option><option value="stress" ${S.demandScenario==='stress'?'selected':''}>供应压力 · 未锁定 +20%</option><option value="low" ${S.demandScenario==='low'?'selected':''}>库存压力 · 未锁定 −20%</option></select></label><div class="seg">${[4,8,13,26].map(n=>`<button data-action="horizon" data-value="${n}" class="${S.horizon===n?'active':''}">${n} 周</button>`).join('')}</div></div>`}
// PAGE IMPLEMENTATIONS FOLLOW
__PAGES_CODE__
// SECTION: DRAWERS / LOCAL WORKFLOW / EVENTS
__INTERACTION_CODE__
render();
// Test surface for calculation and browser validation. No network requests.
window.MemoryDashboard={state:S,data:DATA,tables:T,get models(){return models},get filtered(){return filtered},roll,poAmounts,exposure,monday,render,getEvents};
</script>
</body></html>'''

PAGES_CODE = r'''
function overviewPage(){
 const events=currentEvents(),high=events.filter(r=>(models.find(m=>m.id===r['物料编码'])?.seed['风险等级']==='高'||r['风险事件编号'].startsWith('LOCAL-'))&&r['处置状态']!=='已关闭'),shorts=filtered.filter(m=>m.first>=0&&m.first<S.horizon),urgents=filtered.filter(m=>m.first>=0&&m.first<4),cancelled=filtered.filter(m=>m.cancelled.length),delivery=filtered.filter(m=>m.flags.some(f=>f.id==='SUP-03'));
 const upcoming=events.filter(r=>r['处置状态']!=='已关闭'&&days(r['截止日期'],ASOF)>=0&&days(r['截止日期'],ASOF)<=7).sort((a,b)=>a['截止日期'].localeCompare(b['截止日期']));
 const cards=`<div class="kpis">${moneyKpi('风险关联库存',sum(filtered,m=>m.seed['风险关联库存金额']),'原表汇总 · 按物料去重','accent','来自风险评估结果的模拟值，不代表本页面算法回算。与其他风险金额独立展示，不相加。')}${moneyKpi('战略现货超额 · 估算',sum(filtered,m=>m.strategicExcess),'回算 · 已到货战备归属比例','','仅按已关联批次的台账已到货量 / 现有量分摊现货超额成本。无法实物分离，属于估算。')}${kpi(`未来 ${S.horizon} 周缺料`,fmt(shorts.length),'个料号','回算 · 期末实物库存 < 0','','按每个料号的滚动负库存识别，DRAM 与 NAND 不汇总数量。')}${moneyKpi('不可撤销 PO 承诺',sum(filtered,m=>sum(m.pos,p=>poAmounts(p).locked*num(p['采购单价']))),'回算 · 已考虑变更截止日','','未交量减去仍在变更窗口内的可取消量，再乘采购单价；不假定供应商已同意变更。')}${kpi('高风险未关闭事件',fmt(high.length),'件','事件台账 · 原表风险等级','','原始事件以风险评估结果的原表等级识别；新增本地事件按高风险演示，不与规则命中数混合。')}${kpi('快照日逾期事件',fmt(events.filter(overdueEvent).length),'件',`按 ${ASOF} 判断 · 未关闭`,'','不是以打开文件的当天判断；已缓解仍未关闭。')}</div>`;
 const priority=panel('值得优先关注','从风险信号，找到下一步行动',`<div class="risk-item"><div class="risk-icon">${icon('risks')}</div><div class="body"><h3>4 周内实物缺料</h3><p>先核对承诺到货与可用库存</p></div><strong class="mono">${urgents.length}<small class="muted"> 料</small></strong></div><div class="risk-item"><div class="risk-icon amber">${icon('inventory')}</div><div class="body"><h3>取消项目关联共享料</h3><p>存在库存或 PO，需进一步归因</p></div><strong class="mono">${cancelled.length}<small class="muted"> 料</small></strong></div><div class="risk-item"><div class="risk-icon blue">${icon('supply')}</div><div class="body"><h3>供应商交付异常</h3><p>OTD 低于 ${policy.otd}% 或波动 > 7 天</p></div><strong class="mono">${delivery.length}<small class="muted"> 料</small></strong></div><div style="margin-top:19px"><button class="link" data-action="nav" data-page="risks">进入风险工作台 →</button></div>`);
 const mode=`<div class="seg">${[['balance','库存'],['gap','缺口'],['excess','超目标']].map(([id,label])=>`<button data-action="chart-mode" data-value="${id}" class="${S.chartMode===id?'active':''}">${label}</button>`).join('')}</div>`;
 const top=[...filtered].sort((a,b)=>num(b.seed['风险关联库存金额'])-num(a.seed['风险关联库存金额'])).slice(0,10);
 const cols=[materialCol,col('class','品类',m=>m.mat['DRAM_NAND分类']),col('risk','回算风险',m=>m.level,m=>badge(m.level)),col('stock','风险关联库存 · 原表',m=>num(m.seed['风险关联库存金额']),m=>money(num(m.seed['风险关联库存金额']))),col('cover','覆盖天数',m=>m.coverage,m=>m.coverage===null?'<span class="muted">不适用 · 无需求</span>':`${fmt(m.coverage,1)} 天`),col('trend','13 周需求',m=>sum(m.demand.slice(0,13)),m=>miniBars(m.demand.slice(0,13))),col('driver','主要信号',m=>m.flags[0]?.text||'常规监控',m=>`<span ${tip(m.flags.map(f=>f.id+' · '+f.text).join('\n')||'暂无规则命中')}>${esc(m.flags[0]?.text||'常规监控')}</span>`),col('owner','风险责任人',m=>m.org['风险负责人'])];
 return demandControls()+cards+`<div class="grid-main">${panel(`${S.horizon} 周库存前瞻`,`${S.chartClass} · 同类物料汇总，保留负库存`,rollingChart()+`<div style="display:flex;justify-content:space-between;gap:12px;align-items:center;margin-top:15px">${mode}<span class="note" style="margin:0">${ASOF} 快照 · 周一归桶</span></div>`,chartTools())}${priority}</div><div class="grid-two">${panel('需求与库存，交叉看','同一未来 13 周窗口 · 与前一 MDS 版本比较',bubbleChart())}${panel('接下来的 7 天','以源数据快照为基准的事件截止日',upcoming.length?upcoming.slice(0,4).map(r=>`<div class="risk-item"><div class="risk-icon blue">${icon('clock')}</div><div class="body"><button class="link" data-action="event" data-id="${attr(r['风险事件编号'])}">${esc(r['风险事件编号'])} · ${esc(r['风险描述'])}</button><p>${esc(r['物料编码'])} · ${esc(r['责任人'])}</p></div><span class="tag">${short(r['截止日期'])}</span></div>`).join(''):empty('快照日起 7 天内没有待到期事件。'),`<span class="tag">${upcoming.length} 件</span>`)}</div>`+table('top','重点物料 · TOP 10',cols,top,'按原表风险关联库存金额排序；点击料号查看公式、批次和来源',5);
}
function supplyPage(){
 const atRisk=filtered.filter(m=>m.first>=0&&m.first<S.horizon),ss=filtered.filter(m=>m.firstSafety>=0&&m.firstSafety<S.horizon),delivery=filtered.filter(m=>m.flags.some(f=>f.id==='SUP-03'));
 const cards=`<div class="kpis">${kpi(`${S.horizon} 周实物缺料`,fmt(atRisk.length),'个料号','回算 · 逐料检查负库存','accent')}${kpi('安全库存缺口',fmt(ss.length),'个料号','策略表安全库存只取一次')}${kpi('供应商交付异常',fmt(delivery.length),'个料号','最新统计周期 · OTD / 波动')}${kpi('DRAM 最大累计欠量',qty(sum(filtered.filter(m=>m.mat['DRAM_NAND分类']==='DRAM'),m=>m.maxGap)),'颗','各料号窗口最大负库存之和')}${kpi('NAND 最大累计欠量',qty(sum(filtered.filter(m=>m.mat['DRAM_NAND分类']==='NAND'),m=>m.maxGap)),'颗','各料号窗口最大负库存之和')}${kpi('逾期未交到货',fmt(sum(filtered,m=>m.overdue.length)),'笔','承诺日早于快照 · 未纳入预测')}</div>`;
 const cols=[materialCol,col('plant','库存归属工厂',m=>m.org['工厂']),col('initial','期初净可用 / 颗',m=>m.initial,m=>fmt(m.initial)),col('ss','安全库存 / 颗',m=>m.safety,m=>fmt(m.safety)),col('first','首次实物缺料周',m=>m.first<0?'9999':weeks[m.first],m=>m.first<0?'<span class="positive">26 周内未缺料</span>':`<span class="${m.first<S.horizon?'negative':''}">${weeks[m.first]}${m.first>=S.horizon?' · 窗口外':''}</span>`),col('safety','首次安全缺口周',m=>m.firstSafety<0?'9999':weeks[m.firstSafety],m=>m.firstSafety<0?'—':weeks[m.firstSafety]),col('gap','窗口最大累计欠量 / 颗',m=>m.maxGap,m=>`<span class="${m.maxGap?'negative':''}">${fmt(m.maxGap)}</span>`),col('otd','最新 OTD',m=>num(m.supply?.['OTD']),m=>`<span class="${num(m.supply?.['OTD'])<policy.otd/100?'negative':''}">${pct(num(m.supply?.['OTD']))}</span>`),col('std','交期波动 / 天',m=>num(m.supply?.['交期波动天数'])),col('capacity','产能利用率',m=>num(m.supply?.['产能利用率']),m=>pct(num(m.supply?.['产能利用率']))),col('ppm','质量 PPM',m=>num(m.quality?.['质量PPM']))];
 const chartList=filtered.filter(m=>m.mat['DRAM_NAND分类']===S.chartClass),series=[{name:'周度需求',color:'#0071e3',values:weeks.slice(0,S.horizon).map((_,i)=>sum(chartList,m=>m.demand[i])),area:true},{name:'承诺到货',color:'#7bb59b',values:weeks.slice(0,S.horizon).map((_,i)=>sum(chartList,m=>m.arrivals[i]))}];
 const supplierRows=uniq(filtered.map(m=>m.mat['供应商'])).map(s=>{const all=filtered.filter(m=>m.mat['供应商']===s),total=sum(all,m=>m.supply?.['承诺交货数量']),onTime=sum(all,m=>m.supply?.['实际准时交货数量']);return {supplier:s,count:all.length,otd:total?onTime/total:0,std:sum(all,m=>m.supply?.['交期波动天数'])/all.length}});
 return demandControls()+cards+`<div class="notice">${icon('info')}源 MDS 的周标签为周四，展示统一归到该周周一。第一周沿用源表全周数量，未推断日内消耗。已承诺、未到货 PO 按承诺日归桶；运输中和待发均属于未交量。</div><div class="grid-two">${panel('供应与需求','所选品类独立汇总 · 不合并 DRAM / NAND 数量',chartList.length?lineChart(series,weeks.slice(0,S.horizon)):empty(),chartTools())}${panel('缺口的两种含义','实物不足与安全库存预警，独立识别',rollingChart(filtered,'gap'),`<span class="tag">${S.chartClass}</span>`)}</div>`+table('supply','逐料供应计划',cols,filtered,'负库存继续滚入下一周；最大累计欠量不是逐周缺口相加',8)+table('suppliers','供应商表现',[
 col('supplier','供应商',r=>r.supplier),col('count','所选料号数',r=>r.count),col('otd','加权 OTD',r=>r.otd,r=>pct(r.otd)),col('std','平均交期波动 / 天',r=>r.std,r=>fmt(r.std,1))],supplierRows,'按各料最新周期的承诺量加权 OTD；交期波动为料号简单平均',5);
}
function inventoryPage(){
 const lots=filtered.flatMap(m=>m.lots),oldLots=lots.filter(r=>num(r['库龄天数'])>90),strategies=filtered.flatMap(m=>m.strategies),daily=sum(filtered,m=>m.daily);
 const cards=`<div class="kpis">${moneyKpi('现有库存账面成本',sum(filtered,m=>m.cost),'财务成本表 · 全量持有批次','accent')}${moneyKpi('现货超目标成本',sum(filtered,m=>m.excess*m.unitCost),'回算 · 净可用超出目标部分')}${moneyKpi('战略现货超额 · 估算',sum(filtered,m=>m.strategicExcess),'按已到货战备归属比例')}${moneyKpi('含锁定到货超目标',sum(filtered,m=>m.excessCommitted*m.unitCost),'回算 · 同一库存单位成本估值')}${kpi('超过 90 天库龄',fmt(oldLots.length),'个批次','只看现有批次，不含未来 PO')}${kpi('关联战备决策',fmt(strategies.length),'笔','逐笔保留 PO、批次与退出日期')}</div>`;
 const ageBuckets=[[0,30,'0–30 天','#b7cce3'],[31,60,'31–60 天','#8bafd7'],[61,90,'61–90 天','#4f8bc9'],[91,10000,'90 天以上','#c7a571']].map(([lo,hi,label,color])=>({label,color,value:sum(lots.filter(r=>num(r['库龄天数'])>=lo&&num(r['库龄天数'])<=hi),r=>num(r['现有数量'])*num(r['单位库存成本']))}));
 const totalAge=sum(ageBuckets,r=>r.value),ageChart=`<div class="progress-line">${ageBuckets.map(r=>`<i style="width:${totalAge?r.value/totalAge*100:0}%;background:${r.color}" ${tip(r.label+' · '+money(r.value))}></i>`).join('')}</div>${ageBuckets.map(r=>`<div class="metricrow"><span><i class="dot" style="background:${r.color}"></i>${r.label}</span><strong class="mono">${money(r.value)} <small class="muted">${pct(totalAge?r.value/totalAge:0)}</small></strong></div>`).join('')}`;
 const rawTrend=versions.map(v=>sum(filtered,m=>sum(demandVector(m.id,v,false,false).slice(0,13))*m.unitCost));
 const cols=[materialCol,col('factory','归属工厂',m=>m.org['工厂']),col('onhand','现有量 / 颗',m=>m.onhand,m=>fmt(m.onhand)),col('target','目标量 / 颗',m=>m.target,m=>fmt(m.target)),col('excess','现货超额 / 颗',m=>m.excess,m=>fmt(m.excess)),col('amount','超额成本',m=>m.excess*m.unitCost,m=>money(m.excess*m.unitCost)),col('decline','同窗需求下调',m=>m.decline,m=>m.decline===null?'—':`<span class="${m.decline>0?'amber':''}">${pct(m.decline)}</span>`),col('cover','覆盖 / 天',m=>m.coverage,m=>m.coverage===null?'不适用':fmt(m.coverage,1)),col('lifecycle','生命周期',m=>m.mat['生命周期状态']),col('cancel','取消项目关联',m=>m.cancelled.length,m=>m.cancelled.length?badge('需归因'):'—')];
 const strategyCols=[col('decision','战备决策编号',r=>r['战备决策编号']),col('material','物料',r=>r['物料编码'],r=>materialCell(r['物料编码'])),col('po','对应 PO',r=>r['对应PO号']),col('lot','对应批次',r=>r['对应库存批次']),col('plan','计划量 / 颗',r=>num(r['战备计划数量']),r=>fmt(r['战备计划数量'])),col('arrived','已到货 / 颗',r=>num(r['已到货数量']),r=>fmt(r['已到货数量'])),col('rule','归属规则',r=>r['战略库存归属规则']),col('exit','计划退出',r=>r['计划退出日期']),col('reason','决策原因',r=>r['决策原因'])];
 const ids=new Set(filtered.map(m=>m.id));const substitutes=T['替代料与认证'].filter(r=>ids.has(r['原物料编码'])||ids.has(r['替代物料编码'])).map(r=>{
 const source=models.find(m=>m.id===r['替代物料编码']),target=models.find(m=>m.id===r['原物料编码']),p=PM[r['适用项目编码']],cert=r['认证状态']==='已认证'&&r['技术兼容状态']==='兼容'&&r['认证完成日期']&&r['认证完成日期']<=ASOF,active=p&&!(p['项目状态']==='取消'&&p['取消或暂停日期']<=ASOF),ratio=num(r['替代比例']),output=cert&&active&&source&&target&&ratio>0?Math.min(source.excess,target.maxGap, num(r['可转用数量上限'])/ratio):0;
 return {...r,eligible:cert&&active,upper:output};});
 return demandControls()+cards+`<div class="grid-two">${panel('库存年龄结构','以现有数量 × 批次单位成本展示金额，不跨品类加总颗数',ageChart)}${panel('需求的版本变化','固定未来 13 周窗口 · 原始 MDS 需求的参考成本，未叠加情景',lineChart([{name:'同窗需求参考成本',color:'#0071e3',values:rawTrend,area:true}],versions.map(v=>v.slice(4)),{unit:'万元',scale:1e4}),'<span class="tag">四个版本</span>')}</div><div class="notice warning">${icon('info')}目标库存 = 未来目标天数需求 + 策略安全库存。战备库存采用已到货台账关联批次的比例估算，无法确认归属的数量不计入战略现货超额；取消项目的共享料库存需进一步分摊。</div>`+table('inventory','库存结构与超额识别',cols,filtered,'现货超额不含未交 PO；含锁定到货的超目标规模单列展示，两个金额不可相加',8)+table('strategy','战备决策追溯',strategyCols,strategies,'源数据 · 相同已到货批次的归属量按批次现有量封顶',6)+table('alternatives','替代认证与转用线索',[
 col('original','原物料（需求方）',r=>r['原物料编码'],r=>materialCell(r['原物料编码'])),col('sub','替代物料（供应方）',r=>r['替代物料编码'],r=>materialCell(r['替代物料编码'])),col('project','适用项目',r=>r['适用项目编码']),col('compat','技术兼容',r=>r['技术兼容状态']),col('cert','认证状态',r=>r['认证状态'],r=>badge(r['认证状态'])),col('ratio','替代比例',r=>num(r['替代比例']),r=>pct(num(r['替代比例']))),col('cap','独立测算上限 / 颗',r=>r.upper,r=>fmt(r.upper)),col('state','资格',r=>r.eligible?'可评估':'暂不可用',r=>r.eligible?'<span class="positive">可评估</span>':'<span class="muted">认证 / 项目未满足</span>')],substitutes,'方向：用“替代物料”覆盖“原物料”缺口；各行是独立上限，未执行跨工厂分配，不得相加或视为已转用。',6);
}
function pricePage(){
 const sourcePrices=filtered.flatMap(m=>(idx['价格情景参数'][m.id]||[])),cards=`<div class="kpis">${moneyKpi('情景价格负敞口',sum(filtered,m=>m.priceExposure),`市场价变化 ${S.scenario>0?'+':''}${S.scenario}% · 批次逐一计算`,'accent')}${moneyKpi('当前市场价负敞口',sum(filtered,m=>m.baselineExposure),'价格变动 0% · 同口径基准')}${moneyKpi('现货超额部分敞口 · 估算',sum(filtered,m=>m.excessExposure),'按净可用超额 / 现有量分摊')}${moneyKpi('持有 90 天成本 · 估算',sum(filtered,m=>m.hold),'按当前账面余额恒定估算')}${moneyKpi('现有库存账面成本',sum(filtered,m=>m.cost),'含冻结和已分配的持有批次')}${kpi('可比价格覆盖',fmt(filtered.filter(m=>m.market>0).length),'个物料','CNY · 元/颗 · 最近有效价格')}</div>`;
 const focus=filtered[0],hist=focus?[...focus.prices].reverse():[],presets=[-30,-20,-15,0,15,30],totals=presets.map(v=>sum(filtered,m=>exposure(m.finance,m.market*(1+v/100)))),largest=Math.max(1,...totals),scenarioBars=presets.map((v,i)=>`<div class="metricrow"><span style="min-width:55px">${v>0?'+':''}${v}%</span><div style="flex:1;margin:0 18px"><div class="bar-track" style="width:100%;height:7px"><i style="width:${totals[i]/largest*100}%;background:${v===S.scenario?'#0071e3':'#bdd0e6'}"></i></div></div><strong class="mono" style="font-size:13px">${money(totals[i])}</strong></div>`).join('');
 const priceChart=focus?lineChart([{name:'市场价',color:'#0071e3',values:hist.map(r=>num(r['市场价格']))},{name:'合同价',color:'#95b5a2',values:hist.map(r=>num(r['合同价格'])),dash:true},{name:'当前加权账面成本',color:'#c0aa85',values:hist.map(()=>focus.unitCost),dash:true}],hist.map(r=>r['价格日期']),{unit:'元/颗',scale:1,zero:false,aria:`${focus.id}可比价格历史`}):empty();
 const cols=[materialCol,col('date','最新价格日',m=>m.prices[0]?.['价格日期']),col('cost','加权成本 / 元',m=>m.unitCost,m=>fmt(m.unitCost,2)),col('market','市场价 / 元',m=>m.market,m=>fmt(m.market,2)),col('scenario','情景价 / 元',m=>m.scenarioPrice,m=>fmt(m.scenarioPrice,2)),col('base','基准负敞口',m=>m.baselineExposure,m=>money(m.baselineExposure)),col('exposure','情景负敞口',m=>m.priceExposure,m=>`<span class="${m.priceExposure?'negative':''}">${money(m.priceExposure)}</span>`),col('excess','超额部分 · 估算',m=>m.excessExposure,m=>money(m.excessExposure)),col('hold','90天持有成本',m=>m.hold,m=>money(m.hold))];
 return cards+`<div class="grid-main">${panel('压力测试，轻松推演','相对于各料号最近有效可比市场价 · 批次成本保持不变',`<div class="slider-value mono ${S.scenario<0?'blue':''}" id="scenarioNumber">${S.scenario>0?'+':''}${S.scenario}<span style="font-size:25px;margin-left:4px">%</span></div><label for="scenarioSlider" class="sr-only">市场价格变化百分比</label><input id="scenarioSlider" type="range" min="-30" max="30" step="1" value="${S.scenario}" aria-valuetext="${S.scenario}%"><div style="display:flex;justify-content:space-between;color:var(--muted);font-size:10px;margin-top:7px"><span>−30% 下跌</span><span>0% 基准</span><span>+30% 上涨</span></div><div class="scenario-buttons">${presets.map(v=>`<button class="${S.scenario===v?'active':''}" data-action="scenario" data-value="${v}">${v>0?'+':''}${v}%</button>`).join('')}</div><div class="formula" style="margin-top:23px">负敞口 = Σ 批次数量 × max(0, 批次账面单位成本 − 情景价格)<br>情景价格 = 最近可比市场价 × (1 + 价格变化率)</div><p class="note">情景编号 DEMO-${ASOF.replaceAll('-','')}-${S.scenario<0?'M':'P'}${Math.abs(S.scenario)} · ${RULE_VERSION} · 参数可导出，便于重现。</p><button class="btn" style="margin-top:13px" data-action="export-scenario">${icon('download')}导出情景参数</button>`)}${panel('同口径情景对比','全量持有批次 · CNY',scenarioBars)}</div><div class="notice warning">${icon('info')}价格负敞口、库存成本与风险金额独立展示，不可相加，也不等同于会计减值。90 天持有成本采用当前库存成本恒定假设，包含年化资金成本率与仓储费率。</div>`+table('price','逐料价格敞口',cols,filtered,`情景 ${S.scenario}% · 若价格上涨超过账面成本，敞口最低为 0；每个批次仅计算一次`,8)+`<div class="grid-two">${panel(focus?`${focus.id} · 价格与成本`:'价格与成本','第一条筛选物料；可点击上表任意料号查看独立趋势',priceChart,focus?`<button class="link" data-material="${focus.id}">查看物料 →</button>`:'')}${panel('Excel 中的原始情景','保留原表参数，并与本页实时测算区分',`<p class="paper">原表包含 −30%、−15%、0%、+15%、+30% 五组情景，共 ${sourcePrices.length} 条所选物料记录。滑杆额外支持 −20% 等任意整数百分比，并按最近市场价重新计算。</p><div class="metricrow"><span>当前计算快照</span><strong style="font-size:13px">${ASOF}</strong></div><div class="metricrow"><span>当前情景口径</span><strong style="font-size:13px">批次成本 − 情景市场价</strong></div><button class="btn" data-action="open-source" data-name="价格情景参数" style="margin-top:18px">查看原始情景表 →</button>`)}</div>`;
}
function poPage(){
 const pos=filtered.flatMap(m=>m.pos).map(p=>({...p,_a:poAmounts(p)})).sort((a,b)=>String(b['最后变更日期']).localeCompare(String(a['最后变更日期']))||a['PO号'].localeCompare(b['PO号']));
 const openAmount=sum(pos,r=>r._a.open*r._a.price),cancel=sum(pos,r=>r._a.cancel*r._a.price),defer=sum(pos,r=>r._a.defer*r._a.price),expired=pos.filter(r=>r._a.expired);
 const cards=`<div class="kpis">${moneyKpi('未交 PO 总承诺',openAmount,'未交数量 × 采购单价','accent')}${moneyKpi('窗口内可取消 · 参考',cancel,'取未交量与可取消量较小值')}${moneyKpi('窗口内可延期 · 参考',defer,'与可取消金额不相加')}${moneyKpi('不可撤销承诺',sum(pos,r=>r._a.locked*r._a.price),'截止日已过则视为不可取消')}${kpi('已过最后变更日',fmt(expired.length),'笔',`以快照 ${ASOF} 判断`)}${kpi('窗口内可干预',fmt(pos.filter(r=>r._a.cancel>0||r._a.defer>0).length),'笔','尚需供应商确认与审批')}</div>`;
 const cols=[col('po','PO / 行号',r=>r['PO号'],r=>`<b>${esc(r['PO号'])}</b><br><small class="muted">行 ${esc(r['PO行号'])}</small>`),col('material','物料',r=>r['物料编码'],r=>materialCell(r['物料编码'])),col('supplier','供应商',r=>r['供应商']),col('kind','采购类型',r=>r['采购类型']),col('arrival','承诺到货',r=>r['承诺到货日期']),col('open','未交量 / 颗',r=>r._a.open,r=>fmt(r._a.open)),col('cancel','表列可取消 / 颗',r=>num(r['可取消数量']),r=>fmt(r['可取消数量'])),col('effective','窗口内可取消 / 颗',r=>r._a.cancel,r=>fmt(r._a.cancel)),col('defer','窗口内可延期 / 颗',r=>r._a.defer,r=>fmt(r._a.defer)),col('deadline','最后变更日',r=>r['最后变更日期'],r=>`<span class="${r._a.expired?'negative':''}">${esc(r['最后变更日期'])}</span><br><small class="muted">${r._a.unknown?'日期缺失':r._a.expired?`已过 ${-days(r['最后变更日期'],ASOF)} 天`:`剩余 ${days(r['最后变更日期'],ASOF)} 天`}</small>`),col('locked','不可撤销金额',r=>r._a.locked*r._a.price,r=>money(r._a.locked*r._a.price)),col('fee','原表取消费用',r=>num(r['取消费用']),r=>money(num(r['取消费用']))),col('action','本地处置',r=>r['PO号'],r=>`<button class="link" data-action="new-po-event" data-id="${attr(r['PO号'])}">记录建议 →</button>`)];
 return cards+(expired.length?`<div class="notice warning">${icon('clock')}所选 ${expired.length} 笔 PO 已超过最后变更日。源表的“可取消 / 可延期数量”保留作为合同线索，但不直接计入当前可干预金额。请重新向供应商确认，不以源表配额推断变更仍然有效。</div>`:'')+`<div class="grid-two">${panel('承诺的构成','同一未交数量中的可取消与锁定部分',`<div class="progress-line"><i style="width:${openAmount?cancel/openAmount*100:0}%;background:#8db79d"></i><i style="flex:1;background:#b7cce3"></i></div><div class="metricrow"><span><i class="dot" style="background:#8db79d"></i>窗口内可取消</span><strong>${money(cancel)}</strong></div><div class="metricrow"><span><i class="dot" style="background:#b7cce3"></i>不可撤销承诺</span><strong>${money(openAmount-cancel)}</strong></div><p class="note">“可取消”仍需确认费用与供应商反馈；此处是合同配额的窗口内上限，不是已获批准的节省额。</p>`)}${panel('做出调整之前','所有变更均在本地记录建议',`<div class="audit"><div class="audit-item"><h3>核实变更窗口</h3><p>检查最后变更日及供应商书面确认。</p></div><div class="audit-item"><h3>比较净影响</h3><p>取消费用与可取消金额独立展示；延期不自动记为永久降低风险。</p></div><div class="audit-item"><h3>在事件中留痕</h3><p>记录责任人、措施和证据，本文件不提交 ERP 变更。</p></div></div>`)}</div>`+table('po','未交采购订单',cols,pos,'按最后变更日倒序；运输中、待发均为未交子集，不重复累计',8);
}
function risksPage(){
 const all=currentEvents(),events=all.filter(r=>S.eventStatus==='全部'||S.eventStatus==='逾期'?S.eventStatus!=='逾期'||overdueEvent(r):r['处置状态']===S.eventStatus);
 const cards=`<div class="kpis">${kpi('当前事件',fmt(all.length),'件','原始台账 + 本地新增','accent')}${kpi('处理中',fmt(all.filter(r=>r['处置状态']==='处理中').length),'件','按当前本地状态')}${kpi('待复核 / 已缓解',fmt(all.filter(r=>['待复核','已缓解'].includes(r['处置状态'])).length),'件','已缓解不等同于已关闭')}${kpi('已关闭',fmt(all.filter(r=>r['处置状态']==='已关闭').length),'件','需具备复核金额与证据')}${moneyKpi('源表已验证减少金额',sum(T['风险处置台账'].filter(r=>new Set(all.map(x=>x['风险事件编号'])).has(r['风险事件编号'])),r=>r['已验证减少风险金额']),'模拟台账原值 · 非重新核验')}${kpi('快照日逾期',fmt(all.filter(overdueEvent).length),'件',`截止日 < ${ASOF} 且未关闭`)}</div>`;
 const cols=[col('event','风险事件',r=>r['风险事件编号'],r=>`<button class="link" data-action="event" data-id="${attr(r['风险事件编号'])}">${esc(r['风险事件编号'])} ↗</button><br><small class="muted">${r['风险事件编号'].startsWith('LOCAL-')?'本地新增':'Excel 原始台账'}</small>`),col('material','物料',r=>r['物料编码'],r=>materialCell(r['物料编码'])),col('project','项目',r=>r['项目编码']),col('type','风险类型',r=>r['风险类型']),col('description','风险描述',r=>r['风险描述']),col('initial','初始基线金额',r=>num(r['风险初始金额']),r=>money(num(r['风险初始金额']))),col('owner','责任人',r=>r['责任人']),col('deadline','截止日期',r=>r['截止日期'],r=>`<span class="${overdueEvent(r)?'negative':''}">${esc(r['截止日期'])}</span>`),col('status','状态',r=>r['处置状态'],r=>badge(r['处置状态'])),col('action','操作',r=>r['风险事件编号'],r=>`<button class="link" data-action="edit-event" data-id="${attr(r['风险事件编号'])}">更新进展 →</button>`)];
 const ids=new Set(all.map(r=>r['风险事件编号'])),snapshots=T['风险历史快照'].filter(r=>ids.has(r['风险事件编号'])),dates=uniq(snapshots.map(r=>r['快照日期'])).sort(),amounts=dates.map(d=>sum(snapshots.filter(r=>r['快照日期']===d),r=>r['风险金额']));
 return cards+`<div class="notice">${icon('info')}事件编辑仅保存在当前浏览器的本地工作区。原始 Excel 和历史快照保持只读；建议定期导出工作区备份。无多用户审批、消息通知或服务器审计。</div><div style="display:flex;justify-content:space-between;align-items:center;gap:12px;margin-bottom:16px;flex-wrap:wrap"><div class="seg">${['全部','待确认','处理中','已缓解','待复核','已关闭','逾期'].map(v=>`<button data-action="event-status" data-value="${v}" class="${S.eventStatus===v?'active':''}">${v}</button>`).join('')}</div><div class="actions"><button class="btn" data-action="backup">${icon('download')}工作区备份</button><button class="btn primary" data-action="new-event">${icon('plus')}新建事件</button></div></div>`+table('events','风险事件工作台',cols,events,'点击事件查看历史；转态需填写操作者与原因，关闭需补充复核金额及证据',8)+`<div class="grid-two">${panel('风险历史，原样保留','原始事件快照金额 · 独立情景估算，不与库存金额相加',snapshots.length?lineChart([{name:'台账风险金额',color:'#0071e3',values:amounts,area:true}],dates,{unit:'万元',scale:1e4}):empty())}${panel('本地操作记录','只追加，不覆盖 Excel 的风险历史',local.audit.length?local.audit.slice(-5).reverse().map(r=>`<div class="risk-item"><div class="body"><h3>${esc(r.event)} · ${esc(r.to)}</h3><p>${esc(r.actor)} · ${esc(r.time)}<br>${esc(r.reason)}</p></div></div>`).join(''):'<div class="empty"><strong>还没有本地操作</strong>更新事件后，记录会出现在这里。</div>',`<button class="link" data-action="audit">查看全部 →</button>`)}</div>`;
}
function dataPage(){
 const errors=DATA.quality.filter(q=>q.severity==='error'&&q.failed),warnings=DATA.quality.filter(q=>q.severity==='warning'&&q.failed),policyCount=warnings.reduce((n,r)=>n+r.failed,0);
 const cards=`<div class="sourcecards">${DATA.sources.map((s,i)=>`<article class="panel sourcecard"><span class="tag">SOURCE 0${i+1}</span><div class="num mono">${fmt(s.rows)} <span class="muted" style="font-size:11px">行</span></div><p>${esc(s.name)}<br>${s.sheets.length} 张表 · 快照 ${ASOF}</p><p style="margin-top:9px" ${tip('SHA-256\n'+s.sha256)}>SHA-256 · ${s.sha256.slice(0,14)}…</p></article>`).join('')}</div>`;
 const issues=DATA.quality.map(q=>({...q,status:q.failed?(q.severity==='warning'?'需确认':'异常'):'通过'})),cols=[col('check','校验项',r=>r.name),col('status','结果',r=>r.status,r=>`<span class="tag ${r.status==='通过'?'green':r.status==='需确认'?'amber':'red'}">${r.status}</span>`),col('checked','检查记录数',r=>r.checked,r=>fmt(r.checked)),col('failed','异常 / 差异数',r=>r.failed,r=>r.failed?`<button class="link" data-action="quality-detail" data-name="${attr(r.name)}">${fmt(r.failed)} · 查看</button>`:'0'),col('note','规则说明',r=>r.note)];
 const difference=filtered.map(m=>({id:m.id,seedExcess:num(m.seed['超额库存金额']),calcExcess:m.excess*m.unitCost,seedGap:num(m.seed['缺料数量']),calcGap:Math.max(0,...m.rolling.slice(0,4).map(r=>r.shortage)),seedPrice:num(m.seed['价格负敞口']),calcPrice:m.baselineExposure}));
 const settings=`<div class="settingrow"><div>可用量已扣除分配<p>默认“否”，期初 = 可用 − 已分配；切换只影响本地回算。</p></div><label><input type="checkbox" data-policy="allocatedIncluded" ${policy.allocatedIncluded?'checked':''} aria-label="可用量已扣除分配"></label></div><div class="settingrow"><div>MDS 按物料颗数计<p>默认不重复乘 BOM；关闭后按终端产品台数 × 有效单机用量。</p></div><label><input type="checkbox" data-policy="materialDemand" ${policy.materialDemand?'checked':''} aria-label="MDS按物料颗数计"></label></div><div class="settingrow"><div>剔除取消项目后续需求<p>取消生效周起剔除；暂停项目仍保留原始需求。</p></div><label><input type="checkbox" data-policy="excludeCancelled" ${policy.excludeCancelled?'checked':''} aria-label="剔除取消项目后续需求"></label></div><div class="settingrow"><div>低覆盖预警<p>天数 · 默认 21 天。</p></div><input type="number" min="1" max="365" value="${policy.coverage}" data-policy="coverage" aria-label="低覆盖预警天数"></div><div class="settingrow"><div>供应商 OTD 预警<p>百分比 · 默认 85%。</p></div><input type="number" min="1" max="100" value="${policy.otd}" data-policy="otd" aria-label="OTD预警百分比"></div><div class="settingrow"><div>跌价 20% 敞口预警<p>人民币元 · 初始演示阈值。</p></div><input type="number" min="0" max="1000000000" step="10000" value="${policy.priceLimit}" data-policy="priceLimit" aria-label="跌价敞口预警金额"></div><p class="note">参数仅在本次会话生效，可通过“保存视图”保留并读取。生产规则发布需另行实现审批、版本管理和服务端重算。</p>`;
 const limits=`<div class="metricrow"><span>数据基准</span><strong style="font-size:12px">${ASOF} 固定快照</strong></div><div class="metricrow"><span>致命结构异常</span><strong>${errors.length} <small>项</small></strong></div><div class="metricrow"><span>安全库存口径差异</span><strong>${policyCount} <small>个料号</small></strong></div><p class="paper" style="margin-top:17px">库存归属工厂按组织责任映射；项目需求按共享料池汇总，未虚构跨工厂分配。价格仅采用 CNY / 元每颗的可比参考价。原表风险分和规则权重保留展示，本页没有将供应分与库存分平均抵消。</p><hr class="divider"><h3>已知待确认项</h3><p class="paper" style="margin-top:10px">可用量与已分配量是否重复扣减、MDS 的数量单位、战备已到货的批次唯一归属、共享料池的跨工厂与项目分摊。请在上线前由业务确认。</p><button class="btn" data-action="method" style="margin-top:17px">阅读全部计算口径 →</button>`;
 return cards+`<div class="grid-two">${panel('计算口径，可解释也可切换','所有修改均为本地演示参数',settings)}${panel('数据状态与边界','数据质量不等于算法验收',limits)}</div>`+table('quality','输入质量校验',cols,issues,'生成 HTML 时执行：业务主键、非空外键、PO 守恒、库存与财务勾稽。空值不等同于错误。',8)+table('reconcile','模拟结果与回算对照',[
 col('id','物料',r=>r.id,r=>materialCell(r.id)),col('seedExcess','原表超额金额',r=>r.seedExcess,r=>money(r.seedExcess)),col('calcExcess','回算现货超额',r=>r.calcExcess,r=>money(r.calcExcess)),col('seedGap','原表缺料量 / 颗',r=>r.seedGap,r=>fmt(r.seedGap)),col('calcGap','回算4周最大累计欠量',r=>r.calcGap,r=>fmt(r.calcGap)),col('seedPrice','原表价格负敞口',r=>r.seedPrice,r=>money(r.seedPrice)),col('calcPrice','回算0%负敞口',r=>r.calcPrice,r=>money(r.calcPrice))],difference,'这些值可能采用不同口径，差异不自动判定为错误；模拟结果不能证明回算算法正确。',5)+table('nulls','字段空值概览',[
 col('table','数据表',r=>r['表']),col('field','字段',r=>r['字段']),col('count','空值数',r=>r['空值数']),col('ratio','空值率',r=>r['空值数']/r['总行数'],r=>pct(r['空值数']/r['总行数']))],DATA.nulls,'仅列出存在空值的字段；关闭日期、未关联批次等可为业务允许的空值。',5)+`<div class="panel panel-pad" style="margin-bottom:18px"><div class="panelhead"><div><h2>原始数据浏览器</h2><p>这里保留全部 6,795 行源数据，不受上方业务筛选影响；可使用表内搜索与导出。</p></div></div><select id="sourceSelect" class="data-select" aria-label="选择原始数据表">${Object.keys(T).map(n=>`<option ${S.source===n?'selected':''}>${esc(n)}</option>`).join('')}</select></div>`+rawTable(S.source)+table('rules','Excel 原始规则参数',DATA.tables['风险规则参数'].columns.map(c=>col(c,c,r=>r[c])),T['风险规则参数'],'原表 RULE-xx 为种子配置；本地规则信号按 PRD 的 SUP / INV / PRC / PO 标识执行，不伪称这 7 条原表规则已全部生产化。',7)+`<div class="panel panel-pad"><div class="panelhead"><div><h2>本地工作区</h2><p>事件、操作记录和保存视图均保存在当前浏览器。离线 HTML 不上传任何数据。</p></div></div><div class="actions" style="flex-wrap:wrap"><button class="btn" data-action="backup">导出 JSON 备份</button><button class="btn" data-action="restore">导入工作区备份</button><button class="btn" data-action="reset-local">清除本地演示修改</button><input type="file" id="restoreFile" accept="application/json,.json" hidden></div><p class="note">此导入仅恢复本工具导出的工作区，不解析 Excel。更新四个工作簿后，请运行 code.py 重新生成 HTML。</p></div>`;
}
'''
INTERACTION_CODE = r'''
function drawerHeader(eyebrow,title,sub=''){return `<div class="drawer-head"><div><div class="eyebrow">${esc(eyebrow)}</div><h1 id="drawerTitle">${esc(title)}</h1>${sub?`<p class="subtitle">${esc(sub)}</p>`:''}</div><button class="iconbtn" data-action="close-drawer" aria-label="关闭明细">${icon('close')}</button></div>`}
function openDrawer(type,id=null){if(!S.drawer)returnFocus=document.activeElement;S.drawer={type,id};drawDrawer(true)}
function drawDrawer(focus=false){
 const d=S.drawer;if(!d)return;let body='';
 if(d.type==='material')body=materialDrawer(d.id);
 else if(d.type==='event')body=eventDrawer(d.id);
 else if(d.type==='source')body=drawerHeader('SOURCE DATA',d.id,'保留 Excel 原始行号；不受全局业务筛选影响')+rawTable(d.id,'drawer-source');
 else if(d.type==='method')body=methodDrawer();
 else if(d.type==='audit')body=drawerHeader('LOCAL AUDIT','本地操作记录','原始历史快照不变；本地日志仅用于单人演示')+table('audit','操作记录',[
 col('time','时间（北京时间）',r=>r.time),col('event','事件',r=>r.event),col('actor','操作者',r=>r.actor),col('from','原状态',r=>r.from),col('to','新状态',r=>r.to),col('reason','原因',r=>r.reason),col('evidence','证据说明',r=>r.evidence)],local.audit,'只追加记录 · 可用工作区 JSON 备份保存',12);
 else if(d.type==='quality'){const q=DATA.quality.find(r=>r.name===d.id);body=drawerHeader('DATA QUALITY',q.name,q.note)+`<div class="notice warning">异常 / 差异 ${q.failed} 条。下方最多展示 30 个示例，完整输入可在原始数据浏览器查阅。</div>`+(q.examples.length?table('quality-examples','差异示例',Object.keys(q.examples[0]).map(k=>col(k,k,r=>r[k])),q.examples,'',15):empty('本项校验通过。'))}
 $('#drawerRoot').innerHTML=`<div class="overlay" id="drawerOverlay"><section class="drawer" role="dialog" aria-modal="true" aria-labelledby="drawerTitle" tabindex="-1">${body}</section></div>`;document.body.style.overflow='hidden';if(focus){$('.drawer').scrollTop=0;$('.drawer [data-action="close-drawer"]').focus({preventScroll:true})}
}
function closeDrawer(){S.drawer=null;$('#drawerRoot').innerHTML='';document.body.style.overflow='';if(returnFocus?.isConnected)returnFocus.focus({preventScroll:true});else $('[data-action="nav"][aria-current="page"]')?.focus({preventScroll:true})}
function refresh(opts={}){const opens=$$('details[open][data-filter-details]').map(d=>d.dataset.filterDetails);render();if(S.drawer)drawDrawer();for(const key of opens)$(`[data-filter-details="${key}"]`)?.setAttribute('open','');if(opts.focus){const e=$(opts.focus);if(e){e.focus({preventScroll:true});if(opts.position!=null&&e.setSelectionRange)e.setSelectionRange(opts.position,opts.position)}}}
function materialDrawer(id){
 const m=models.find(r=>r.id===id);if(!m)return drawerHeader('MATERIAL DETAIL','物料不存在')+empty();
 const weeklyCols=[col('week','周一日期',r=>r.week),col('opening','期初净可用',r=>r.opening,r=>fmt(r.opening)),col('arrival','承诺到货',r=>r.arrival,r=>fmt(r.arrival)),col('demand','物料需求',r=>r.demand,r=>fmt(r.demand)),col('balance','期末实物库存',r=>r.balance,r=>`<span class="${r.balance<0?'negative':''}">${fmt(r.balance)}</span>`),col('shortage','实物缺料',r=>r.shortage,r=>fmt(r.shortage)),col('safetyGap','安全库存缺口',r=>r.safetyGap,r=>fmt(r.safetyGap))];
 const sourceLotCols=[col('lot','批次',r=>r['库存批次']),col('row','Excel 行',r=>r.__row),col('factory','工厂',r=>r['工厂']),col('onhand','现有',r=>r['现有数量'],r=>fmt(r['现有数量'])),col('available','可用',r=>r['可用数量'],r=>fmt(r['可用数量'])),col('frozen','冻结',r=>r['冻结数量'],r=>fmt(r['冻结数量'])),col('allocated','已分配',r=>r['已分配数量'],r=>fmt(r['已分配数量'])),col('cost','单位成本',r=>r['单位库存成本'],r=>fmt(r['单位库存成本'],2)),col('age','库龄 / 天',r=>r['库龄天数'])];
 const projectList=m.relatedProjects.map(p=>`${p['项目编码']} · ${p['项目名称']}（${p['项目状态']}）`).join('；');
 const sourcePOCols=[col('po','PO',r=>r['PO号']),col('date','承诺到货',r=>r['承诺到货日期']),col('quantity','未交数量',r=>num(r['未交数量']),r=>fmt(r['未交数量'])),col('deadline','最后变更',r=>r['最后变更日期']),col('type','采购类型',r=>r['采购类型'])];
 const pricing=m.prices.slice().reverse();
 return drawerHeader('MATERIAL DETAIL',`${m.id} · ${m.mat['物料名称']}`,`${m.mat['DRAM_NAND分类']} / ${m.org['工厂']} / ${m.mat['供应商']}`)+`<div class="kpis">${kpi('净可用库存',qty(m.initial),'颗','已按当前分配口径扣减')}${kpi('覆盖天数',m.coverage===null?'—':fmt(m.coverage,1),m.coverage===null?'无需求':'天','未来 28 天平均需求')}${moneyKpi('情景价格负敞口',m.priceExposure,`当前情景 ${S.scenario}%`)}</div>`+panel('回算过程',`${S.version} · ${RULE_VERSION}`,`<div class="formula">期初净可用 = ${fmt(m.available)}${policy.allocatedIncluded?'（可用量已扣分配）':` − ${fmt(m.allocated)}（已分配）`} = ${fmt(m.initial)} 颗<br>目标库存 = ${fmt(m.target-m.safety)}（${m.targetDays} 天需求） + ${fmt(m.safety)}（安全库存） = ${fmt(m.target)} 颗<br>现货超额 = max(0, ${fmt(m.initial)} − ${fmt(m.target)}) = ${fmt(m.excess)} 颗<br>现货超额成本 = ${fmt(m.excess)} × ${fmt(m.unitCost,4)} = ${money(m.excess*m.unitCost)}<br>情景价格 = ${fmt(m.market,2)} × (1 ${S.scenario<0?'−':'+'} ${Math.abs(S.scenario)}%) = ${fmt(m.scenarioPrice,4)} 元/颗</div><p class="note">覆盖天数为净可用 / 未来四周日均需求；无需求时显示不适用。批次账面成本取财务成本表，不以单一均价代替批次负敞口。</p>`)+panel('26 周逐周库存','单位：颗 · 只看当前物料，不跨品类汇总',lineChart([{name:'期末库存',color:'#0071e3',values:m.rolling.map(r=>r.balance),area:true},{name:'安全库存',color:'#c4c4ca',values:weeks.map(()=>m.safety),dash:true}],weeks,{height:235}))+table('material-weeks','逐周计算明细',weeklyCols,m.rolling,'I(t) = I(t−1) + 到货 − 需求；负值继续结转，不将每周缺口相加',8)+panel('规则与归属','规则信号是本地回算；原表分数单独保留',`<div style="display:flex;gap:6px;flex-wrap:wrap">${m.flags.length?m.flags.map(f=>`<span class="tag ${f.level==='严重'?'red':'amber'}">${esc(f.id)} · ${esc(f.text)}</span>`).join(''):'<span class="tag green">未命中当前演示规则</span>'}</div><dl class="definition" style="margin-top:20px"><dt>关联项目</dt><dd>${esc(projectList)}</dd><dt>库存归属</dt><dd>${esc(m.org['工厂'])} · ${esc(m.org['业务单元'])} · ${esc(m.org['产品族'])}</dd><dt>风险责任人</dt><dd>${esc(m.org['风险负责人'])}</dd><dt>原表供应 / 库存分</dt><dd>${fmt(num(m.seed['供应风险分']))} / ${fmt(num(m.seed['库存风险分']))}，各自独立</dd><dt>原表主要驱动</dt><dd>${esc(m.seed['主要风险驱动'])}</dd><dt>原表 / 回算等级</dt><dd>${badge(m.seed['风险等级'])} / ${badge(m.level)}</dd><dt>战备已到货归属</dt><dd>${fmt(m.stratQty)} 颗，按关联批次封顶</dd><dt>首次缺料周</dt><dd>${m.first<0?'26 周内未缺料':weeks[m.first]}</dd></dl><p class="note">INV-02 表示取消项目关联的共享料池存在库存或 PO，未证明整池库存都因取消项目而形成风险。</p>`)+table('material-lots','库存批次与来源',sourceLotCols,m.lots,'来源：02_需求计划与库存供应_模拟数据.xlsx / 库存批次快照',5)+table('material-pos','未交 PO 与来源',sourcePOCols,m.pos,'来源：02_需求计划与库存供应_模拟数据.xlsx / 采购订单明细',5)+panel('价格历史','CNY / 元每颗 · 只展示可比口径',lineChart([{name:'市场价',color:'#0071e3',values:pricing.map(r=>num(r['市场价格']))},{name:'加权账面成本',color:'#bdab8b',values:pricing.map(()=>m.unitCost),dash:true}],pricing.map(r=>r['价格日期']),{unit:'元/颗',scale:1,zero:false,height:210}))+`<p class="note">run_id: DEMO-${ASOF}-${S.version}-${RULE_VERSION} · source_batch_id: ${DATA.sources[1].sha256.slice(0,12)}<br>所有事实来自内嵌工作簿快照；外部 Excel 更新后请重新生成本文件。</p>`;
}
function methodDrawer(){
 const entries=[
 ['01 · 时间与版本',`数据时点固定为 ${ASOF}。选择截至该时点已发布的 MDS。历史版本切换只替换需求，库存、PO、价格仍使用 ${ASOF} 快照，不等同于历史完整回放。源表周四标签统一归到周一；第一周不做无依据的日级拆分。`],
 ['02 · 期初、需求与到货','默认期初净可用 = Σ可用 − Σ已分配；可在数据配置切换“已扣分配”。默认 MDS 为物料需求，不重复乘 BOM；切换终端产品需求时按有效 BOM 用量换算。取消项目生效后剔除，暂停项目保留原值。未锁定需求可按 ±20% 演示情景调整。PO 按承诺到货日归桶，运输中和待发不额外叠加。逾期未交的到货日期未知，不擅自挪到第一周。'],
 ['03 · 工厂与项目','工厂、业务单元、产品族筛选采用组织责任映射，是物料共享池的归属维度。源数据没有完备的跨工厂分配账，因此不把项目工厂伪装为库存工厂。项目筛选限定需求与事件，并选择关联物料；库存、PO 保留这些物料的完整共享池。'],
 ['04 · 逐周滚动与缺料','I(0)=净可用；I(t)=I(t−1)+A(t)−D(t)。保留负值结转。实物缺料=max(0,−I(t))；安全库存缺口=max(0,SS−I(t))。SS 按物料+工厂取最近有效策略一次。最大累计欠量取窗口内负库存绝对值的最大值，不把各周缺口相加。DRAM/NAND 数量分开展示。'],
 ['05 · 覆盖与超额','覆盖天数 = 净可用 /（未来4周需求 / 28），需求为零显示不适用。目标库存 = 目标天数内需求 + 策略安全库存，按周需求比例覆盖非整周天数。现货超额 = max(0,净可用−目标)。含锁定到货的超目标量 = max(0,净可用+不可取消未交量−目标)，按当前加权账面成本参考估值。两者独立展示。'],
 ['06 · 战备归因与替代','战略现货超额按“战备台账已到货、已关联批次数量 / 现有量”分摊现货超额成本，批次归属量以该批次现有量封顶，属于虚拟比例估算。空批次不视为已到货库存。替代方向为：替代物料提供供应、原物料存在缺口；仅兼容、已认证、认证已生效且项目未取消时计算独立上限 min(替代料现货超额,原物料窗口缺口,可转用上限/替代比例)。没有多对多占用分配，不能叠加上限。'],
 ['07 · 价格与持有成本','价格采用最近有效的 CNY / 元每颗现货参考价；负敞口 = Σ批次数量×max(0,批次账面单位成本−情景价格)，逐批次计算。超额部分按超额量/现有量比例分摊，标记为估算。持有90天成本 = Σ当前账面余额×（年化资金率+仓储费率）×90/365，是恒定库存近似。所有敞口不代表会计减值，不与库存或事件金额相加。'],
 ['08 · MDS 变化与规则','需求下调率 =（上版−本版）/上版，比较相同未来13周，使用原始需求、不叠加 ±20% 情景与取消剔除；上版为零时不适用。SUP-01、SUP-02、SUP-03、INV-01、INV-02、INV-03、PRC-01、PO-01 为可解释演示信号。INV-03 按13周无总需求近似，未实现确定性分层。未声称完整事件去重、自动升级或生产规则权重系统。原表风险分保留，不与回算混合。'],
 ['09 · PO 窗口','有效可取消量 = 窗口内 min(未交量,可取消量)，窗口截止日包含当天。截止日缺失或已过则保守置零。不可撤销金额 =（未交量−有效可取消量）×单价；延期配额另列，不与取消配额相加。源表配额是参考上限，不表示已获得供应商批准；取消费用不自动扣减或确认为节省。'],
 ['10 · 事件、关闭与本地存储','事件可新增、指派、更新措施、转态、驳回与重开。源状态“已缓解”保留，不能自动等同关闭。每次修改记录操作者、时间、前后状态和原因。关闭要求复核金额、复核日期、证据及归因；复核减少额=初始基线−同口径复核金额，延期不能归因于永久消除。本地记录可由用户修改或清除，不是服务器可信审计；不提供多用户 RBAC、审批、附件上传、通知或 ERP 操作。'],
 ['11 · 数据质量与复现','HTML 生成时检查21张表业务主键、非空外键、PO守恒、库存和财务勾稽；空值概览不把业务可空字段判为错误。各源文件保留 SHA-256，表中保留 Excel 行号。原表风险评估只是合成种子数据，与本页回算的差异另列，不用其证明算法正确。更新工作簿后使用 code.py 重新生成；所有处理均在本地。']
 ];
 return drawerHeader('METHODOLOGY','让口径先于结论。','可解释的本地分析基线 · PRD V1.0')+entries.map(([title,text])=>panel(esc(title),'',`<p class="paper">${esc(text)}</p>`)).join('');
}
function eventDrawer(id){
 const e=getEvents().find(r=>r['风险事件编号']===id);if(!e)return drawerHeader('RISK DETAIL','事件不存在')+empty();
 const history=T['风险历史快照'].filter(r=>r['风险事件编号']===id),audit=local.audit.filter(r=>r.event===id),review=e['本地复核金额'];
 return drawerHeader('RISK DETAIL',id,e['风险描述'])+`<div class="kpis">${moneyKpi('初始基线金额',num(e['风险初始金额']),'事件自身的情景口径')}${moneyKpi('源表已验证减少',num(T['风险处置台账'].find(r=>r['风险事件编号']===id)?.['已验证减少风险金额']),'模拟台账原值，未重新核验')}${kpi('当前状态',esc(e['处置状态']),'',e['风险事件编号'].startsWith('LOCAL-')?'本地新增事件':'原始事件 + 本地修改')}</div>`+panel('事件档案','详情与当前进展',`<dl class="definition"><dt>物料</dt><dd>${materialCell(e['物料编码'])}</dd><dt>项目</dt><dd>${esc(e['项目编码'])}</dd><dt>风险类型</dt><dd>${esc(e['风险类型'])}</dd><dt>责任人 / 协作人</dt><dd>${esc(e['责任人'])} / ${esc(e['协作人']||'—')}</dd><dt>截止日期</dt><dd>${esc(e['截止日期'])} ${overdueEvent(e)?badge('逾期'):''}</dd><dt>建议动作</dt><dd>${esc(e['建议动作'])}</dd><dt>确认动作</dt><dd>${esc(e['确认动作']||'尚未确认')}</dd><dt>本地证据</dt><dd>${esc(e['本地证据']||'尚未填写')}</dd><dt>复核归因</dt><dd>${esc(e['本地归因']||'尚未填写')}</dd><dt>本地复核金额</dt><dd>${review!==undefined?money(num(review)):'尚未复核'}</dd><dt>本地复核减少额</dt><dd>${review!==undefined?money(num(e['风险初始金额'])-num(review)):'尚未复核'}</dd></dl><p class="note">本地复核与源表的“已验证减少金额”分开保存；金额口径应由复核人确认，延期不得自动记录为永久消除。</p><button class="btn primary" data-action="edit-event" data-id="${attr(id)}" style="margin-top:18px">更新进展</button>`)+panel('原始风险历史','保持 Excel 快照，不以本地编辑覆盖',history.length?lineChart([{name:'事件风险金额',color:'#0071e3',values:history.map(r=>num(r['风险金额'])),area:true}],history.map(r=>r['快照日期']),{unit:'万元',scale:1e4,height:200}):empty('本地新增事件暂无 Excel 历史快照。'))+panel('本地操作轨迹','按追加时间排序',audit.length?`<div class="audit">${audit.map(a=>`<div class="audit-item"><h3>${esc(a.from||'新建')} → ${esc(a.to)}</h3><p>${esc(a.actor)} · ${esc(a.time)}<br>${esc(a.reason)}${a.evidence?'<br>证据：'+esc(a.evidence):''}</p></div>`).join('')}</div>`:empty('尚无本地修改。'));
}
const TRANSITIONS={'新发现':['新发现','待确认'],'待确认':['待确认','已指派','处理中'],'已指派':['已指派','处理中','待确认'],'处理中':['处理中','待复核'],'已缓解':['已缓解','处理中','待复核'],'待复核':['待复核','处理中','已关闭'],'已关闭':['已关闭','处理中']};
function dialogContent(title,body,footer=''){const d=$('#dialog');d.innerHTML=`<div class="dialog-head"><h2 id="dialogTitle">${esc(title)}</h2><button class="iconbtn" data-action="close-dialog" aria-label="关闭对话框">${icon('close')}</button></div>${body}${footer}`;d.setAttribute('aria-labelledby','dialogTitle');if(!d.open)d.showModal();return d}
function editEvent(id){
 const r=getEvents().find(x=>x['风险事件编号']===id);if(!r)return;
 const states=TRANSITIONS[r['处置状态']]||[r['处置状态'],'待确认','处理中'];
 dialogContent(`更新 ${id}`,`<form id="eventForm" data-id="${attr(id)}"><div class="fieldgrid"><label class="field">状态<select name="status" id="editStatus">${states.map(v=>`<option ${v===r['处置状态']?'selected':''}>${esc(v)}</option>`).join('')}</select></label><label class="field">责任人<input name="owner" required maxlength="60" value="${attr(r['责任人'])}"></label><label class="field">截止日期<input name="deadline" type="date" required value="${attr(r['截止日期'])}"></label><label class="field">操作者<input name="actor" required maxlength="60" placeholder="填写真实或演示姓名"></label><label class="field full">确认动作<textarea name="action" maxlength="1000" placeholder="记录下一步措施">${esc(r['确认动作']||'')}</textarea></label><label class="field full">本次修改原因 · 必填<textarea name="reason" maxlength="1000" required placeholder="说明指派、转态、驳回或重开的原因"></textarea></label><label class="field">同口径复核金额 / 元<input name="review" type="number" min="0" step="0.01" value="${r['本地复核金额']??''}" placeholder="关闭时必填"></label><label class="field">复核日期<input name="reviewDate" type="date" value="${attr(r['复核日期']||ASOF)}"></label><label class="field">效果归因<select name="attribution"><option value="">请选择</option>${['已确认消除','暂时转移','自然变化'].map(v=>`<option ${v===r['本地归因']?'selected':''}>${v}</option>`).join('')}</select></label><label class="field">是否涉及延期 PO<select name="deferred"><option value="no" ${r['本地涉及延期']?'':'selected'}>否</option><option value="yes" ${r['本地涉及延期']?'selected':''}>是 · 不记为永久消除</option></select></label><label class="field full">证据 / 复核说明<textarea name="evidence" maxlength="2000" placeholder="填写验证方法、证据编号或文档链接；本工具不上传附件">${esc(r['本地证据']||'')}</textarea></label></div><p class="note">可保存进展而保持原状态。关闭要求：确认动作、复核金额、复核日期、证据与归因。已关闭事件可重开到处理中；已缓解需先进入待复核。</p><p class="negative small" id="formError" role="alert" style="margin-top:10px"></p><div class="dialog-foot"><button type="button" class="btn" data-action="close-dialog">取消</button><button type="submit" class="btn primary">保存到本地</button></div></form>`);
}
function newEvent(prefill={}){
 const id=prefill.material||filtered[0]?.id||MATERIALS[0]['物料编码'];
 dialogContent('新建风险事件',`<form id="newEventForm"><div class="fieldgrid"><label class="field">物料<select name="material" id="newMaterial">${MATERIALS.map(m=>`<option value="${m['物料编码']}" ${m['物料编码']===id?'selected':''}>${m['物料编码']} · ${esc(m['物料名称'])}</option>`).join('')}</select></label><label class="field">项目<select name="project" id="newProject">${eventProjectOptions(id,prefill.project)}</select></label><label class="field">风险类型<select name="type">${['缺料','砍单','超储','跌价','交付延迟','PO锁定'].map(t=>`<option ${prefill.po&&t==='PO锁定'?'selected':''}>${t}</option>`).join('')}</select></label><label class="field">初始基线金额 / 元<input type="number" min="0" step="0.01" required name="amount" value="${prefill.amount??''}" placeholder="按事件口径填写"></label><label class="field full">风险描述<input name="description" required maxlength="200" value="${attr(prefill.po?'PO '+prefill.po+' 的变更评估':'')}"></label><label class="field">责任人<input name="owner" required maxlength="60" value="${attr(ORG[id]?.['风险负责人']||'')}"></label><label class="field">截止日期<input type="date" name="deadline" required value="${iso(utc(ASOF)+7*MS)}"></label><label class="field full">建议动作<textarea name="action" required maxlength="1000" placeholder="措施、目标与确认方式"></textarea></label><label class="field">操作者<input name="actor" required maxlength="60"></label><label class="field">创建原因<input name="reason" required maxlength="300"></label></div><p class="note">新增事件采用待确认状态，编号以 LOCAL 开头，不写回 Excel。事件金额不能与库存成本和其他情景敞口相加。</p><p class="negative small" id="formError" role="alert" style="margin-top:10px"></p><div class="dialog-foot"><button type="button" class="btn" data-action="close-dialog">取消</button><button type="submit" class="btn primary">创建事件</button></div></form>`);
}
function eventProjectOptions(id,selected=''){const ids=uniq((projectLinks[id]||[]).map(r=>r['项目编码']));return '<option value="">共享池 / 不指定项目</option>'+ids.map(p=>`<option value="${p}" ${p===selected?'selected':''}>${p} · ${esc(PM[p]?.['项目名称'])}</option>`).join('')}
function saveEvent(form){
 const f=Object.fromEntries(new FormData(form)),id=form.dataset.id,r=getEvents().find(r=>r['风险事件编号']===id),error=$('#formError');
 const fail=s=>{error.textContent=s;return false};
 if(!f.actor.trim()||!f.reason.trim()||!f.owner.trim())return fail('操作者、责任人和修改原因不能只包含空格。');
 if(!(TRANSITIONS[r['处置状态']]||[r['处置状态'],'待确认','处理中']).includes(f.status))return fail('不允许的状态转换。');
 if(f.status==='已指派'&&!f.owner.trim())return fail('指派必须指定责任人。');
 const closing=f.status==='已关闭';if(closing&&(!f.action.trim()||f.review===''||!f.reviewDate||!f.evidence.trim()||!f.attribution))return fail('关闭事件需要确认动作、复核金额、复核日期、证据和效果归因。');
 if(f.review!==''&&(!Number.isFinite(Number(f.review))||Number(f.review)<0))return fail('复核金额须为非负数字。');
 if(closing&&f.deferred==='yes'&&f.attribution==='已确认消除')return fail('延期 PO 不能归因为永久消除，请选择“暂时转移”或“自然变化”。');
 const patch={'处置状态':f.status,'责任人':f.owner.trim(),'截止日期':f.deadline,'确认动作':f.action.trim(),'本地证据':f.evidence.trim(),'本地归因':f.attribution,'本地涉及延期':f.deferred==='yes'};
 if(f.review!=='')patch['本地复核金额']=Number(f.review);if(f.reviewDate)patch['复核日期']=f.reviewDate;if(closing)patch['关闭日期']=f.reviewDate;else if(r['处置状态']==='已关闭')patch['关闭日期']=null;
 local.events[id]={...(local.events[id]||{}),...patch};local.audit.push({event:id,actor:f.actor.trim(),time:nowCN(),from:r['处置状态'],to:f.status,reason:f.reason.trim(),evidence:f.evidence.trim(),before:{...r},after:{...r,...patch}});persist();$('#dialog').close();refresh();toast('进展已记录到本地工作区');
}
function createEvent(form){
 const f=Object.fromEntries(new FormData(form));if(!f.description.trim()||!f.action.trim()||!f.actor.trim()||!f.owner.trim()||!f.reason.trim()){ $('#formError').textContent='必填文本不能只包含空格。';return}
 if(!MM[f.material]||f.project&&!(projectLinks[f.material]||[]).some(x=>x['项目编码']===f.project)){$('#formError').textContent='所选物料与项目没有有效关联。';return}
 const id='LOCAL-'+Date.now().toString(36).toUpperCase()+'-'+Math.random().toString(36).slice(2,5).toUpperCase(),r={'风险事件编号':id,'发现日期':ASOF,'物料编码':f.material,'项目编码':f.project||null,'风险类型':f.type,'风险描述':f.description.trim(),'风险初始金额':Number(f.amount),'建议动作':f.action.trim(),'确认动作':null,'责任人':f.owner.trim(),'协作人':null,'截止日期':f.deadline,'处置状态':'待确认','预计减少风险金额':0,'已验证减少风险金额':0,'复核日期':null,'关闭日期':null,'备注':'本地演示事件；基准日采用源数据快照，操作时间见日志'};
 local.created.push(r);local.audit.push({event:id,actor:f.actor.trim(),time:nowCN(),from:'',to:'待确认',reason:f.reason.trim(),evidence:'',before:null,after:{...r}});persist();$('#dialog').close();refresh();toast('本地事件已创建');openDrawer('event',id);
}
function csvCell(v){if(v===null||v===undefined)return '""';let text=String(v);if(typeof v!=='number'&&/^\s*[=+@\-\t\r]/.test(text))text="'"+text;return '"'+text.replaceAll('"','""')+'"'}
function download(name,body,type){const blob=new Blob([body],{type}),url=URL.createObjectURL(blob),a=document.createElement('a');a.href=url;a.download=name;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(url),2000)}
function filterSummary(){return {as_of_date:ASOF,version:S.version,horizon:S.horizon,filters:S.filters,query:S.query,demand_scenario:S.demandScenario,price_scenario_percent:S.scenario,rule_version:RULE_VERSION,policy:{...policy}}}
function exportTable(id){const t=tableRegistry[id];if(!t)return;const meta=[['Memory 内存战略风险管理 · 合成演示数据'],['导出时间（北京时间）',nowCN()],['数据快照',ASOF],['筛选 / 口径',JSON.stringify(filterSummary())],['表名',t.title,'范围',id.includes('source')?'原始全表（仅应用表内搜索）':'当前业务筛选与表内搜索'],[],t.cols.map(c=>c.label)];const rows=[...meta,...t.rows.map(r=>t.cols.map(c=>c.value(r)))];download(t.title.replace(/[<>:"/\\|?*]/g,'_')+'.csv','\ufeff'+rows.map(r=>r.map(csvCell).join(',')).join('\r\n'),'text/csv;charset=utf-8');toast(`已导出 ${t.rows.length} 条记录，包含筛选和快照信息`)}
function backup(){download('Memory-工作区-'+ASOF+'.json',JSON.stringify({format:'memory-workspace',version:1,asof:ASOF,exported_at:nowCN(),local},null,2),'application/json');toast('工作区备份已导出')}
function validateBackup(v){
 if(!v||v.format!=='memory-workspace'||v.version!==1||v.asof!==ASOF||!v.local||typeof v.local!=='object'||Array.isArray(v.local))throw Error('备份格式或快照日期不匹配');
 const x=v.local;if(!Array.isArray(x.created)||!Array.isArray(x.audit)||!Array.isArray(x.views)||!x.events||typeof x.events!=='object'||Array.isArray(x.events))throw Error('备份字段不完整');
 const sourceIDs=new Set(T['风险处置台账'].map(r=>r['风险事件编号'])),ids=new Set(sourceIDs);for(const r of x.created){if(!r||typeof r['风险事件编号']!=='string'||!r['风险事件编号'].startsWith('LOCAL-')||ids.has(r['风险事件编号'])||!MM[r['物料编码']]||!TRANSITIONS[r['处置状态']]||!Number.isFinite(r['风险初始金额'])||r['风险初始金额']<0||!/^\d{4}-\d{2}-\d{2}$/.test(r['截止日期']||'')||typeof r['风险描述']!=='string'||typeof r['责任人']!=='string')throw Error('新增事件内容不合法');if(r['项目编码']&&!(projectLinks[r['物料编码']]||[]).some(p=>p['项目编码']===r['项目编码']))throw Error('新增事件项目引用不合法');ids.add(r['风险事件编号'])}
 const allowed=new Set(['处置状态','责任人','截止日期','确认动作','本地证据','本地归因','本地涉及延期','本地复核金额','复核日期','关闭日期']);for(const [id,p]of Object.entries(x.events)){if(!ids.has(id)||!p||typeof p!=='object'||Object.keys(p).some(k=>!allowed.has(k))||p['处置状态']&&!TRANSITIONS[p['处置状态']])throw Error('事件修改字段不合法');if(p['本地复核金额']!==undefined&&(!Number.isFinite(p['本地复核金额'])||p['本地复核金额']<0))throw Error('复核金额不合法');if(p['截止日期']&&!/^\d{4}-\d{2}-\d{2}$/.test(p['截止日期']))throw Error('截止日格式不合法');if(p['处置状态']==='已关闭'&&(!p['本地证据']||p['本地复核金额']===undefined||!p['本地归因']||!p['确认动作']||!p['复核日期']||p['本地涉及延期']&&p['本地归因']==='已确认消除'))throw Error('关闭事件缺少有效复核信息')}
 if(x.audit.some(a=>!a||!ids.has(a.event)||typeof a.actor!=='string'||typeof a.reason!=='string'||typeof a.time!=='string'))throw Error('操作记录不合法');for(const view of x.views){if(!view||typeof view.name!=='string'||!versions.includes(view.version)||![4,8,13,26].includes(view.horizon)||!view.filters||Object.keys(view.filters).some(k=>!FILTERS[k]||!Array.isArray(view.filters[k])||view.filters[k].some(v=>!FILTERS[k][1]().includes(v))))throw Error('保存视图不合法');validatePolicy(view.policy||{})}
 return {events:x.events,created:x.created,audit:x.audit,views:x.views};
}
function validatePolicy(p){for(const k of ['allocatedIncluded','materialDemand','excludeCancelled'])if(p[k]!==undefined&&typeof p[k]!=='boolean')throw Error('口径参数格式错误');for(const [key,min,max] of [['coverage',1,365],['otd',1,100],['priceLimit',0,1e9]])if(p[key]!==undefined&&(!Number.isFinite(p[key])||p[key]<min||p[key]>max))throw Error('口径参数超出范围')}
function saveView(){dialogContent('保存当前视图',`<form id="saveViewForm"><label class="field">视图名称<input required name="name" maxlength="40" placeholder="例如：武汉工厂 · 13周供应"></label><p class="note">保存全局筛选、MDS 版本、窗口、需求/价格情景与计算口径，不保存临时表格页码。</p><div class="dialog-foot"><button type="button" class="btn" data-action="close-dialog">取消</button><button class="btn primary">保存视图</button></div></form>`)}
function loadView(){dialogContent('读取已保存视图',local.views.length?local.views.map((v,i)=>`<div class="risk-item"><div class="body"><h3>${esc(v.name)}</h3><p>${esc(v.version)} · ${v.horizon} 周 · ${esc(v.time||'')}</p></div><button class="btn" data-action="apply-view" data-id="${i}">读取</button></div>`).join(''):empty('尚无保存的视图。'))}
function resetFilters(){for(const k of Object.keys(S.filters))S.filters[k]=[];S.query='';S.tables={};refresh()}
function livePrice(){const temp=document.createElement('div');temp.innerHTML=pricePage();for(const selector of ['.kpis','#table-price','.grid-main > .panel:last-child']){const old=$('#main '+selector),next=temp.querySelector(selector);if(old&&next)old.replaceWith(next)}const n=$('#scenarioNumber');if(n)n.innerHTML=`${S.scenario>0?'+':''}${S.scenario}<span style="font-size:25px;margin-left:4px">%</span>`;const a=$('#main .grid-main > .panel:first-child .note'),b=temp.querySelector('.grid-main > .panel:first-child .note');if(a&&b)a.innerHTML=b.innerHTML;$$('[data-action="scenario"]').forEach(el=>el.classList.toggle('active',num(el.dataset.value)===S.scenario));$('#scenarioSlider')?.setAttribute('aria-valuetext',S.scenario+'%')}
document.addEventListener('submit',e=>{
 const f=e.target;if(f.id==='eventForm'){e.preventDefault();saveEvent(f)}else if(f.id==='newEventForm'){e.preventDefault();createEvent(f)}else if(f.id==='saveViewForm'){e.preventDefault();const name=String(new FormData(f).get('name')).trim();if(!name)return;const view={name,time:nowCN(),...structuredClone(filterSummary()),page:S.page};local.views=local.views.filter(v=>v.name!==name);local.views.push(view);if(local.views.length>20)local.views.shift();persist();$('#dialog').close();refresh();toast('当前筛选与口径已保存')}
});
document.addEventListener('click',e=>{
 if(e.target.id==='drawerOverlay'){closeDrawer();return}
 if(!e.target.closest('.filter'))$$('details.filter[open]').forEach(d=>d.removeAttribute('open'));
 const material=e.target.closest('[data-material]');if(material){openDrawer('material',material.dataset.material);return}
 const b=e.target.closest('[data-action]');if(!b)return;const a=b.dataset.action,id=b.dataset.id,v=b.dataset.value;
 if(a==='nav'){if(S.drawer)closeDrawer();S.page=b.dataset.page;$('#sidebar').classList.remove('open');refresh();window.scrollTo({top:0,behavior:'instant'})}
 else if(a==='menu')$('#sidebar').classList.toggle('open');
 else if(a==='method')openDrawer('method');else if(a==='print')window.print();else if(a==='close-drawer')closeDrawer();else if(a==='close-dialog')$('#dialog').close();
 else if(a==='reset')resetFilters();else if(a==='remove-filter'){S.filters[b.dataset.key]=S.filters[b.dataset.key].filter(x=>x!==v);refresh()}
 else if(a==='chart-class'){S.chartClass=v;refresh()}else if(a==='chart-mode'){S.chartMode=v;refresh()}else if(a==='horizon'){S.horizon=num(v);refresh()}else if(a==='scenario'){S.scenario=num(v);refresh()}
 else if(a==='sort'){const st=S.tables[id];st.dir=st.sort===b.dataset.key?-st.dir:1;st.sort=b.dataset.key;st.page=1;refresh()}
 else if(a==='paginate'){S.tables[id].page+=num(b.dataset.step);const y=window.scrollY,dy=$('.drawer')?.scrollTop;refresh();window.scrollTo(0,y);if(dy!==undefined&&$('.drawer'))$('.drawer').scrollTop=dy}
 else if(a==='export')exportTable(id);else if(a==='event')openDrawer('event',id);else if(a==='edit-event')editEvent(id);else if(a==='new-event')newEvent();
 else if(a==='new-po-event'){const po=T['采购订单明细'].find(r=>r['PO号']===id);newEvent({material:po['物料编码'],project:po['项目编码'],po:id,amount:poAmounts(po).locked*num(po['采购单价'])})}
 else if(a==='event-status'){S.eventStatus=v;S.tables.events&&(S.tables.events.page=1);refresh()}
 else if(a==='open-source')openDrawer('source',b.dataset.name);else if(a==='quality-detail')openDrawer('quality',b.dataset.name);else if(a==='audit')openDrawer('audit');
 else if(a==='backup')backup();else if(a==='restore')$('#restoreFile').click();
 else if(a==='reset-local'){if(window.confirm('清除本地新增事件、修改、日志和保存视图？原始 Excel 与内嵌数据不会被修改。')){local={events:{},created:[],audit:[],views:[]};persist();refresh();toast('本地演示修改已清除')}}
 else if(a==='save-view')saveView();else if(a==='load-view')loadView();else if(a==='apply-view'){const view=local.views[num(id)];if(view){for(const k of Object.keys(S.filters))S.filters[k]=[...(view.filters[k]||[])];S.query=view.query||'';S.version=view.version;S.horizon=view.horizon;S.demandScenario=view.demand_scenario||'base';S.scenario=view.price_scenario_percent??-20;S.tables={};if(view.policy)Object.assign(policy,view.policy);$('#dialog').close();refresh();toast('已读取视图：'+view.name)}}
 else if(a==='export-scenario'){download(`Memory-情景-${S.scenario}pct.json`,JSON.stringify({...filterSummary(),run_id:`DEMO-${ASOF}-${S.version}-${RULE_VERSION}`,source_hashes:DATA.sources.map(s=>({name:s.name,sha256:s.sha256})),results:filtered.map(m=>({material:m.id,market_price:m.market,scenario_price:m.scenarioPrice,exposure:m.priceExposure,excess_exposure_estimate:m.excessExposure}))},null,2),'application/json');toast('情景参数与逐料结果已导出')}
});
document.addEventListener('change',async e=>{
 const t=e.target;
 if(t.dataset.filter){const key=t.dataset.filter;if(t.checked)S.filters[key]=uniq([...S.filters[key],t.value]);else S.filters[key]=S.filters[key].filter(v=>v!==t.value);S.tables={};refresh()}
 else if(t.dataset.column){const st=S.tables[t.dataset.table],key=t.dataset.column;if(t.checked)st.hidden=st.hidden.filter(k=>k!==key);else st.hidden=uniq([...st.hidden,key]);refresh()}
 else if(t.id==='versionSelect'){S.version=t.value;refresh()}
 else if(t.id==='demandScenario'){S.demandScenario=t.value;refresh()}
 else if(t.id==='scenarioSlider'){S.scenario=num(t.value);refresh()}
 else if(t.id==='sourceSelect'){S.source=t.value;refresh()}
 else if(t.dataset.policy){const key=t.dataset.policy,value=t.type==='checkbox'?t.checked:num(t.value);try{validatePolicy({[key]:value});policy[key]=value;refresh();toast('已按新的演示口径重新计算')}catch(err){toast(err.message);refresh()}}
 else if(t.id==='newMaterial'){$('#newProject').innerHTML=eventProjectOptions(t.value);$('#newEventForm [name="owner"]').value=ORG[t.value]?.['风险负责人']||''}
 else if(t.id==='restoreFile'){const file=t.files[0];if(!file)return;if(file.size>8*1024*1024){toast('备份超过 8 MB 上限');return}try{const v=validateBackup(JSON.parse(await file.text()));if(!window.confirm('导入备份将替换当前本地工作区，是否继续？建议先导出当前备份。'))return;local=v;persist();refresh();toast('工作区已恢复')}catch(err){toast('无法导入：'+err.message)}}
});
let searchTimer;
document.addEventListener('input',e=>{
 const t=e.target;if(t.dataset.optionSearch){const q=t.value.toLowerCase();t.closest('.popover').querySelectorAll('[data-option-label]').forEach(l=>l.hidden=!l.dataset.optionLabel.toLowerCase().includes(q))}
 else if(t.id==='globalSearch'){clearTimeout(searchTimer);const value=t.value,pos=t.selectionStart;searchTimer=setTimeout(()=>{S.query=value;S.tables={};refresh({focus:'#globalSearch',position:pos})},160)}
 else if(t.dataset.tableSearch){clearTimeout(searchTimer);const id=t.dataset.tableSearch,value=t.value,pos=t.selectionStart;searchTimer=setTimeout(()=>{S.tables[id].search=value;S.tables[id].page=1;refresh({focus:`[data-table-search="${id}"]`,position:pos})},160)}
 else if(t.id==='scenarioSlider'){S.scenario=num(t.value);buildModels();filtered=models.filter(matches);livePrice()}
});
document.addEventListener('keydown',e=>{
 if(e.key==='Escape'){if($('#dialog').open)return;if(S.drawer){e.preventDefault();closeDrawer()}else $('#sidebar').classList.remove('open');$('#tooltip').style.display='none'}
 if((e.key==='Enter'||e.key===' ')&&e.target.matches('svg [data-material]')){e.preventDefault();openDrawer('material',e.target.dataset.material)}
 if(e.key==='Tab'&&S.drawer&&!$('#dialog').open){const items=$$('.drawer button:not(:disabled),.drawer a,.drawer input,.drawer select,.drawer textarea,.drawer [tabindex="0"],.drawer summary').filter(x=>x.getClientRects().length);const first=items[0],last=items.at(-1);if(e.shiftKey&&document.activeElement===first){e.preventDefault();last?.focus()}else if(!e.shiftKey&&document.activeElement===last){e.preventDefault();first?.focus()}}
});
document.addEventListener('pointerover',e=>{const t=e.target.closest('[data-tip]');if(t){const box=$('#tooltip');box.textContent=t.dataset.tip;box.style.display='block'}});
document.addEventListener('pointermove',e=>{const b=$('#tooltip');if(b.style.display==='block'){b.style.left=Math.max(8,Math.min(e.clientX+14,window.innerWidth-b.offsetWidth-10))+'px';b.style.top=Math.max(8,Math.min(e.clientY+14,window.innerHeight-b.offsetHeight-10))+'px'}});
document.addEventListener('pointerout',e=>{if(e.target.closest('[data-tip]'))$('#tooltip').style.display='none'});
'''
README = r'''# Memory · 内存战略备货风险管理看板

基于《看版prd.docx》V1.0 与四个 Excel 工作簿制作的离线交互看板。采用 Apple 风格的系统字体、浅灰背景、大留白、细边框与蓝色交互强调。

## 三个交付文件

| 文件 | 用途 |
| --- | --- |
| `看板.html` | 可独立使用的完整看板，已内嵌全部数据、CSS、JavaScript 和 SVG 图表 |
| `README.md` | 使用方式、功能、业务口径、边界和验证方式 |
| `code.py` | 完整可复现源码：Excel 加载与校验、数据嵌入、HTML/CSS/JavaScript、本文档生成 |

看板使用时只需 HTML 文件，不需要 Python、服务器、网络、字体下载或 CDN。源码重新生成时，需要原始四个 Excel 文件。

## 打开与操作

1. 将 `看板.html` 下载到电脑，用 Chrome、Edge 或 Safari 的近期版本直接打开；不建议在微信、邮件的受限预览器中运行。
2. 通过左侧导航切换七个页面。窄屏设备可使用左上角菜单。
3. 全局筛选支持业务单元、工厂、产品族、项目、物料、供应商、DRAM/NAND、回算风险方向和等级，多选条件在各页联动。不同维度取交集，同一维度取并集。可搜索物料、关联项目或责任人，使用“保存视图 / 读取视图”保留筛选、版本、情景和口径。
4. 点击蓝色料号或风险矩阵气泡，查看逐周公式、期初库存、到货、需求、安全库存、批次、来源行号和价格历史。
5. 表格支持表内搜索、点击表头排序、列显隐、分页与 CSV 导出。导出包含当前筛选、快照、规则口径和北京时间，并导出所有匹配记录而不只当前页。为保持可追溯性，CSV 保留全部列。UTF-8 BOM 便于 Excel 打开，文本类公式前缀会被转义。
6. 价格滑杆支持 −30%～+30%，拖动时同步更新敞口、情景对比及逐料结果。可导出完整情景参数 JSON。
7. 风险处置页支持本地新建、指派、措施确认、进展记录、待复核、关闭与重开。操作者和原因必填。关闭还需确认动作、同口径复核金额、复核日期、证据及效果归因。
8. 使用“工作区备份”导出事件修改、操作日志、保存视图。数据质量页可恢复本工具同一快照的 JSON 备份，导入时会要求确认替换。浏览器清理数据、换浏览器或移动 HTML 路径可能使本地工作区无法继续读取，建议及时备份。
9. 右上角“打印 / PDF”可打印当前页面。打印的是当前页面当前表格页码；完整明细请用 CSV 导出。

如浏览器限制 `file://` 本地存储，可在 HTML 所在目录运行 `python3 -m http.server 8000` 并通过本机浏览器访问。此命令是可选的本地运行方式，HTML 本身没有联网请求。

## 七个页面

| 页面 | 主要内容 |
| --- | --- |
| 管理总览 | 六个 KPI、4/8/13/26 周库存前瞻、缺口与超目标切换、需求下调 × 覆盖月数气泡矩阵、重点物料、七天到期事件 |
| 供应保障 | 净可用、实物缺料、安全库存缺口、供应需求趋势、首次缺料周、最新 OTD、交期波动、产能利用率与质量 PPM |
| 战略库存 | 现货与含锁定到货的超目标测算、战略归属比例估算、库龄金额、同窗 MDS 版本变化、战备决策与替代认证线索 |
| 价格与情景 | 批次价格压力测试、0% 基准负敞口、超额部分估算、90 天持有成本、可比价格与合同价格趋势 |
| PO 干预 | 未交承诺、表列配额与窗口内有效配额、取消与延期独立展示、不可撤销金额、截止日、费用与本地建议记录 |
| 风险处置 | 原始事件与本地进展、状态筛选、责任人、截止日、证据、同口径复核、Excel 原始风险历史与本地操作轨迹 |
| 数据质量与配置 | 四个数据源与 SHA-256、输入校验、口径切换、模拟风险与回算对照、空值率、全部21张原始表浏览、规则参数、工作区备份 |

## 数据来源

固定业务时点：**2026-10-22**。原始内容为合成演示数据，不能当作当前真实经营状态。

| Excel | 工作表数 | 数据行数 |
| --- | ---: | ---: |
| `01_物料与项目基础数据_模拟数据.xlsx` | 5 | 114 |
| `02_需求计划与库存供应_模拟数据.xlsx` | 6 | 5,648 |
| `03_战备成本与价格风险_模拟数据.xlsx` | 5 | 923 |
| `04_风险模型与处置闭环_模拟数据.xlsx` | 5 | 110 |
| 合计 | 21 | 6,795 |

包含25个物料、12个项目、50个库存批次、75笔PO、18条初始风险事件、54条风险历史快照。四个 MDS 版本为 2026-10-01、10-08、10-15、10-22。

输入更新后应重新运行生成器；直接打开现有 HTML 不会读取旁边变化的 Excel。原始数据浏览器明确展示完整源表，不受上方业务筛选影响，仅应用该表自己的搜索条件。其余业务页按全局条件联动。

## 核心计算与假设

- **时间基准**：所有经营截止日、逾期与价格取值使用 Excel 快照日；本地操作时间和导出时间使用北京时间。历史 MDS 切换仅替换需求，库存、PO、价格仍为当前固定快照，不是历史全量回放。
- **周口径**：源 MDS 的周标签为周四，图表归入该自然周的周一，第一桶为2026-10-19。源数据未提供周内消耗分布，保留首周完整需求，不伪造逐日精度。
- **期初净可用**：默认 `Σ可用数量 − Σ已分配数量`。可在数据质量页切换“可用量已扣分配”。源表的可用+冻结能与现有勾稽，但分配是否已经扣减仍需业务确认。
- **需求**：默认 MDS 数量是物料颗数，不重复乘 BOM。关闭该选项后按有效 BOM 用量换算。取消项目生效后需求剔除；暂停项目保留原始需求。±20% 需求情景只改变未锁定需求。
- **范围归属**：工厂、业务单元、产品族依照组织责任映射，属于物料共享池维度。项目筛选限定需求与事件，并选择关联物料；库存与PO保留这些物料的完整共享池。数据没有完备的跨工厂分配账，不能把它当作项目专属库存或工厂独立 ATP 计划。
- **到货去重**：未交PO按承诺日期归桶，运输中与待发是未交子集，不重复相加。早于快照日仍未交的PO由于缺乏新的承诺日期，不擅自计入首周到货，并在供应页计数提示。
- **滚动库存**：`I(t)=I(t−1)+A(t)−D(t)`，保留负值结转。实物缺料=`max(0,−I(t))`，安全库存缺口=`max(0,SS−I(t))`。最大累计欠量是窗口内负库存绝对值的最大值，不能把连续周缺口相加。
- **安全库存**：只按物料+工厂取最近有效库存策略一次，批次安全库存不重复扣。现有快照与策略的安全库存差异会单独列出。
- **覆盖**：净可用库存 / 未来28天日均需求。需求为0显示“不适用”，不除以0。
- **目标与现货超额**：目标量=未来目标天数需求+安全库存，非整周按比例取需求。现货超额=`max(0,净可用−目标)`。超额成本用财务表加权单位成本估值。
- **含锁定到货超目标**：`max(0,净可用+不可取消未交量−目标)`，采用当前库存单位成本参考估值；与现货超额独立展示，不相加。
- **战略现货归因**：仅对台账中已到货且关联批次的数量进行虚拟归属，各批次封顶为其现有量；再按归属量/现有量分摊现货超额成本。标注“估算”，未证明物理隔离，也不把全部共享库存当成战略库存。
- **MDS 下调**：`(上版需求−本版需求)/上版需求`，比较固定未来13周的原始需求，不叠加取消剔除或±20%情景。上版为0时不适用；连续下调需要三个版本。
- **负敞口**：`Σ批次数量×max(0,批次账面单位成本−最近市场价×(1+变动率))`。每个持有批次只算一次，包含冻结/已分配的实物持有量；超额部分按比例估算。价格限定 CNY / 元每颗。
- **持有成本**：`当前账面成本×(年化资金成本率+仓储费率)×90/365`，假定当前余额恒定，属于近似。
- **PO 有效窗口**：截至最后变更日当天允许使用配额。窗口内有效可取消量=`min(未交量,可取消量)`，过期或日期未知保守取0。延期单独展示，不与取消额相加。金额为参考上限，仍需供应商确认与审批。
- **替代线索**：用替代物料超额覆盖原物料缺口。要求兼容、认证完成、生效日期有效且项目未取消；独立上限=`min(替代料现货超额,原物料窗口最大缺口,可转用上限/替代比例)`。没有多对多占用分配和跨工厂物流约束，不得叠加独立上限。
- **金额独立**：风险关联库存、超额成本、负敞口、事件风险金额互不相加，不等于财务减值。

### 原始模拟值与回算的区别

“原表”标识表示直接来自种子数据，特别是风险关联库存、原始供应/库存分、原始已验证减少金额和风险历史。页面回算不假装能还原所有原始模拟结果，数据质量页提供并排对照。

高风险事件 KPI 采用原始风险等级映射到事件；新增本地事件按高风险演示。回算风险等级/方向筛选按PRD演示信号：SUP-01/02/03、INV-01/02/03、PRC-01、PO-01。严重信号优先，再高风险；未命中显示低风险。它不是完整权重模型，不将供应分与库存分平均抵消。

INV-02 表示取消项目关联的共享料池仍有库存或承诺，未完成项目独占归因。INV-03 采用13周无总需求近似，未实现确定性分层与生产事件合并。

原表全部75笔PO的最后变更日为2026-10-17，早于快照10月22日，因此窗口内有效取消与延期为0。该结果是数据事实，不是页面缺失；表列合同配额仍可查阅。

### 事件工作流

主流程：待确认 → 已指派 / 处理中 → 待复核 → 已关闭。已关闭可以重开到处理中；待复核可驳回到处理中。原表“已缓解”保留，后续可进入处理中或待复核，不能自动关闭。

每次保存追加本地日志，保留前后记录、操作者、时间、原因和证据。关闭时输入复核金额后，本地复核减少额为初始基线减去同口径复核金额；源表“已验证减少金额”单独保留。涉及延期PO时不能选择“已确认消除”。

## 重新生成

环境已提供 Python 3.12.14 与 openpyxl 3.1.5。生成器仅依赖 openpyxl；前端为原生JavaScript和SVG，不需 npm。

```bash
cd /workspace/data-panel-sc
python3 code.py --check
python3 code.py
```

其他机器可以先使用隔离环境：

```bash
python3 -m venv .venv
.venv/bin/python -m pip install openpyxl==3.1.5
.venv/bin/python code.py --data-dir /path/to/source --output /path/to/output
```

输入目录应包含原始四个工作簿，建议也包含原始 PRD，生成器会记录其 SHA-256。`code.py`、HTML 与README均以UTF-8保存；输出默认位于源码所在目录，重新生成会覆盖该目录内的 `看板.html` 和 `README.md`。不修改任何原始 Excel 或 PRD。若指定其他输出目录，源码仍保留在运行时的位置。

`--check` 不写输出，检查21张表主键、非空物料/项目/PO/批次/事件引用、PO数量守恒、库存勾稽与财务勾稽。结构性错误会以非零退出；安全库存口径差异属于需确认警告，可在看板中查阅。

源码结构：`load_data` 负责 Excel 读取与输入质量检查；`HTML` 包含界面、样式、状态及纯计算函数；`PAGES_CODE` 包含七个业务页面；`INTERACTION_CODE` 包含表格、抽屉、事件工作流、备份和导出交互；`README` 为本文内容。

## 验证与使用边界

已对实际四个工作簿运行生成器校验，并通过 33 项 Chromium 自动化检查，覆盖七页渲染、筛选、物料下钻、价格滑杆、版本/窗口切换、排序/分页/列控制、导出、保存视图、事件创建/流转/关闭校验、刷新恢复、备份校验、断网运行及390/768像素窄屏布局。计算边界检查覆盖负库存结转、零需求、PO截止日、批次敞口封顶及分类数量分离；全部25个物料的26周回算结果另与独立的 Python 参考计算核对。

本交付为可离线使用的分析与流程演示版本。**不包括**生产数据库、API、ERP接入、Excel在线导入、多用户登录/RBAC、可信服务端审计、审批权限、附件上传、通知调度、自动风险解除与去重、百万行扩展目标或会计减值核算。这些需要在后续后端工程中实现，不能由三个静态交付文件提供。默认阈值与业务假设尚需业务确认，不能直接用作正式采购与财务决策。

HTML 内含所有原始模拟数据；拿到文件的人可以查看全部内容。界面筛选不构成权限隔离。请按数据共享范围分发，并在进入真实数据阶段另行设计鉴权与脱敏。
'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--check", action="store_true", help="Validate inputs without writing outputs")
    args = parser.parse_args()
    data = load_data(args.data_dir)
    failures = [q for q in data["quality"] if q["severity"] == "error" and q["failed"]]
    print(f"Loaded {len(data['sources'])} workbooks, {len(data['tables'])} sheets, {data['rowCount']:,} rows; as of {data['asof']}")
    for issue in data["quality"]:
        if issue["failed"]:
            print(f"{issue['severity'].upper()}: {issue['name']}: {issue['failed']}")
    if failures:
        raise SystemExit("Input integrity validation failed. Source workbooks were not modified; see diagnostics above.")
    if args.check:
        print("PASS: business keys, foreign keys, PO quantities, stock and finance reconciliation")
        return
    embedded = json.dumps(data, ensure_ascii=False, separators=(",", ":"), allow_nan=False).replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    html = HTML.replace("__PAGES_CODE__", PAGES_CODE).replace("__INTERACTION_CODE__", INTERACTION_CODE).replace("__DASHBOARD_DATA__", embedded)
    args.output.mkdir(parents=True, exist_ok=True)
    path = args.output / "看板.html"
    path.write_text(html, encoding="utf-8")
    (args.output / "README.md").write_text(README, encoding="utf-8")
    print(f"Wrote {path} ({path.stat().st_size:,} bytes) and README.md")


if __name__ == "__main__":
    main()
